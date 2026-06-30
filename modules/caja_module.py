"""
Módulo de Caja Diaria - Club Los Pocitos Azufrados
Apertura, cierre y movimientos de caja diaria
"""

import tkinter as tk
from tkinter import ttk, messagebox
import os, sys, datetime, logging

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE_DIR)

from utils.tema_corporativo import COLORES, FUENTES, crear_boton, format_money, aplicar_estilo_tabla
from database.connection import get_connection, conexion_segura
from utils.logger import log_auditoria

_log = logging.getLogger("pocitos")


class CajaModule:
    def __init__(self, parent, usuario, callback_ir_bar=None):
        self.parent = parent
        self.usuario = usuario
        self.caja_actual = None
        self.movimientos = []
        self._callback_ir_bar = callback_ir_bar

        self._crear_interfaz()
        self._cargar_caja_actual()

    def _crear_interfaz(self):
        """Interfaz principal de caja"""
        # Header
        header = tk.Frame(self.parent, bg=COLORES['primario'])
        header.pack(fill='x')

        inner_header = tk.Frame(header, bg=COLORES['primario'], padx=15, pady=10)
        inner_header.pack(fill='x')

        tk.Label(inner_header, text="Caja Diaria",
                 font=FUENTES['encabezado'], fg=COLORES['texto_claro'],
                 bg=COLORES['primario']).pack(side='left')

        self.lbl_estado = tk.Label(inner_header, text="",
                                   font=FUENTES['encabezado'], fg=COLORES['acento'],
                                   bg=COLORES['primario'])
        self.lbl_estado.pack(side='right')

        # ========== CUERPO PRINCIPAL ==========
        body = tk.Frame(self.parent, bg=COLORES['fondo'])
        body.pack(fill='both', expand=True)

        # Canvas scrollable
        self._canvas_caja = tk.Canvas(body, bg=COLORES['fondo'], highlightthickness=0)
        sb = ttk.Scrollbar(body, orient='vertical', command=self._canvas_caja.yview)
        self._canvas_caja.configure(yscrollcommand=sb.set)
        sb.pack(side='right', fill='y')
        self._canvas_caja.pack(side='left', fill='both', expand=True)

        scroll_frame = tk.Frame(self._canvas_caja, bg=COLORES['fondo'])
        self._win = self._canvas_caja.create_window((0, 0), window=scroll_frame, anchor='nw')
        self._canvas_caja.bind('<Configure>', lambda e: self._canvas_caja.itemconfig(self._win, width=e.width))
        scroll_frame.bind('<Configure>', lambda e: self._canvas_caja.configure(scrollregion=self._canvas_caja.bbox('all')))

        # MouseWheel binding
        def _on_wheel(e):
            self._canvas_caja.yview_scroll(int(-1*(e.delta/120)), "units")
        self._canvas_caja.bind_all('<MouseWheel>', _on_wheel)
        self._canvas_caja.bind_all('<Button-4>', lambda e: self._canvas_caja.yview_scroll(-1, 'units'))
        self._canvas_caja.bind_all('<Button-5>', lambda e: self._canvas_caja.yview_scroll(1, 'units'))

        # Agregar padding interior
        body_inner = tk.Frame(scroll_frame, bg=COLORES['fondo'])
        body_inner.pack(fill='both', expand=True, padx=15, pady=15)
        body = body_inner

        # ========== SECCIÓN SUPERIOR: ESTADO ACTUAL ==========
        estado_frame = tk.Frame(body, bg=COLORES['fondo_card'],
                               highlightbackground=COLORES['borde'], highlightthickness=1)
        estado_frame.pack(fill='x', padx=0, pady=(0, 15))

        estado_inner = tk.Frame(estado_frame, bg=COLORES['fondo_card'], padx=20, pady=15)
        estado_inner.pack(fill='x')

        # Fila 1: Estado y Monto Inicial
        fila1 = tk.Frame(estado_inner, bg=COLORES['fondo_card'])
        fila1.pack(fill='x', pady=5)

        tk.Label(fila1, text="Estado:", font=FUENTES['encabezado'],
                 fg=COLORES['texto'], bg=COLORES['fondo_card']).pack(side='left')
        self.lbl_estado_text = tk.Label(fila1, text="CERRADA",
                                        font=FUENTES['encabezado'], fg=COLORES['error'],
                                        bg=COLORES['fondo_card'])
        self.lbl_estado_text.pack(side='left', padx=10)

        tk.Label(fila1, text="Monto Inicial:", font=FUENTES['normal'],
                 fg=COLORES['texto_secundario'], bg=COLORES['fondo_card']).pack(side='left', padx=(50, 0))
        self.lbl_inicial = tk.Label(fila1, text="$ 0",
                                    font=FUENTES['precio'], fg=COLORES['primario'],
                                    bg=COLORES['fondo_card'])
        self.lbl_inicial.pack(side='left', padx=10)

        # Botones de apertura/cierre
        btn_frame = tk.Frame(estado_inner, bg=COLORES['fondo_card'])
        btn_frame.pack(fill='x', pady=(10, 0))

        self.btn_apertura = crear_boton(btn_frame, "Abrir Caja", self._abrir_caja, tipo='exito')
        self.btn_apertura.pack(side='left', padx=2)

        self.btn_reporte_x = crear_boton(btn_frame, "Reporte X (Turno)", self._reporte_x, tipo='acento')
        self.btn_reporte_x.pack(side='left', padx=2)

        self.btn_cierre = crear_boton(btn_frame, "Reporte Z (Cierre del Dia)", self._cerrar_caja, tipo='error')
        self.btn_cierre.pack(side='left', padx=2)

        self.btn_bar = crear_boton(btn_frame, "Ir al Bar", lambda: self._ir_al_bar(), tipo='primario')
        self.btn_bar.pack(side='left', padx=2)

        self.btn_gasto_rapido = crear_boton(btn_frame, "Gasto Rapido", self._registrar_gasto_rapido, tipo='secundario')
        self.btn_gasto_rapido.pack(side='left', padx=2)

        # ========== SECCIÓN MEDIA: RESUMEN ==========
        resumen_frame = tk.Frame(body, bg=COLORES['fondo'])
        resumen_frame.pack(fill='x', padx=0, pady=(0, 15))

        for i, (label, var_name) in enumerate([
            ("Total Ventas", 'total_ventas'),
            ("Cuentas Abiertas", 'cuentas_abiertas'),
            ("Total Gastos", 'total_gastos'),
            ("Esperado", 'esperado')
        ]):
            card = tk.Frame(resumen_frame, bg=COLORES['fondo_card'],
                           highlightbackground=COLORES['borde'], highlightthickness=1)
            card.pack(side='left', fill='both', expand=True, padx=(0 if i == 0 else 8))

            card_inner = tk.Frame(card, bg=COLORES['fondo_card'], padx=15, pady=12)
            card_inner.pack(fill='both', expand=True)

            tk.Label(card_inner, text=label, font=FUENTES['pequena'],
                     fg=COLORES['texto_secundario'], bg=COLORES['fondo_card']).pack(anchor='w')

            lbl_valor = tk.Label(card_inner, text="$ 0", font=FUENTES['precio'],
                                fg=COLORES['primario'], bg=COLORES['fondo_card'])
            lbl_valor.pack(anchor='w', pady=(5, 0))

            setattr(self, f'lbl_{var_name}', lbl_valor)

        # ========== SECCIÓN INFERIOR: MOVIMIENTOS ==========
        mov_label = tk.Label(body, text="Movimientos de Hoy", font=FUENTES['subtitulo'],
                             fg=COLORES['texto'], bg=COLORES['fondo'])
        mov_label.pack(anchor='w', pady=(10, 8))

        tabla_frame = tk.Frame(body, bg=COLORES['fondo_card'],
                               highlightbackground=COLORES['borde'], highlightthickness=1)
        tabla_frame.pack(fill='both', expand=True)

        self.tree_movimientos = ttk.Treeview(tabla_frame, height=10,
                                             columns=('tipo', 'concepto', 'valor', 'metodo', 'usuario', 'hora'),
                                             show='headings')

        self.tree_movimientos.column('tipo', width=80, anchor='center')
        self.tree_movimientos.column('concepto', width=200, anchor='w')
        self.tree_movimientos.column('valor', width=100, anchor='e')
        self.tree_movimientos.column('metodo', width=100, anchor='center')
        self.tree_movimientos.column('usuario', width=100, anchor='center')
        self.tree_movimientos.column('hora', width=100, anchor='center')

        self.tree_movimientos.heading('tipo', text='Tipo')
        self.tree_movimientos.heading('concepto', text='Concepto')
        self.tree_movimientos.heading('valor', text='Valor')
        self.tree_movimientos.heading('metodo', text='Método')
        self.tree_movimientos.heading('usuario', text='Usuario')
        self.tree_movimientos.heading('hora', text='Hora')

        aplicar_estilo_tabla(self.tree_movimientos)
        self.tree_movimientos.pack(fill='both', expand=True)

    def _cargar_caja_actual(self):
        """Carga la caja de hoy"""
        try:
            with conexion_segura() as conn:
                hoy = datetime.date.today().isoformat()

                caja = conn.execute("""
                    SELECT * FROM caja_diaria
                    WHERE DATE(fecha_apertura) = ?
                """, (hoy,)).fetchone()

                self.caja_actual = caja

                if caja and caja['estado'] == 'ABIERTA':
                    # Caja abierta
                    estado_text = "ABIERTA"
                    color = COLORES['exito']
                    self.lbl_estado_text.config(text=estado_text, fg=color)
                    self.lbl_inicial.config(text=format_money(caja['monto_inicial']))

                    # Total ventas: consultar desde ventas (no el campo estático de caja_diaria)
                    hoy_str = datetime.date.today().isoformat()
                    total_ventas_real = conn.execute("""
                        SELECT COALESCE(SUM(total), 0) FROM ventas
                        WHERE DATE(fecha_creacion) = ? AND estado IN ('PAGADA', 'CERRADA')
                    """, (hoy_str,)).fetchone()[0]
                    self.lbl_total_ventas.config(text=format_money(total_ventas_real))

                    # Calcular cuentas abiertas (sum of saldo_pendiente WHERE estado='ABIERTA')
                    cuentas = conn.execute("""
                        SELECT SUM(saldo_pendiente) FROM ventas
                        WHERE estado='ABIERTA'
                    """).fetchone()[0] or 0
                    self.lbl_cuentas_abiertas.config(text=format_money(cuentas))

                    self.lbl_total_gastos.config(text=format_money(caja['total_gastos']))
                    self.lbl_esperado.config(text=format_money(caja['monto_esperado']))

                    # Botones
                    self.btn_apertura.config(state='disabled')
                    self.btn_reporte_x.config(state='normal')
                    self.btn_cierre.config(state='normal')

                    # Cargar movimientos
                    movs = conn.execute("""
                        SELECT tipo, concepto, valor, metodo_pago, usuario, fecha_hora
                        FROM movimientos_caja
                        WHERE id_caja = ?
                        ORDER BY fecha_hora DESC
                    """, (caja['id_caja'],)).fetchall()

                    for item in self.tree_movimientos.get_children():
                        self.tree_movimientos.delete(item)

                    for i, mov in enumerate(movs):
                        tag = 'alerta' if mov['tipo'] == 'EGRESO' else 'par' if i % 2 == 0 else 'impar'
                        self.tree_movimientos.insert('', 'end', values=(
                            mov['tipo'],
                            mov['concepto'],
                            format_money(mov['valor']),
                            mov['metodo_pago'],
                            mov['usuario'],
                            mov['fecha_hora'][11:16]
                        ), tags=(tag,))

                else:
                    # Sin caja abierta (o cerrada)
                    cerrada_hoy = caja is not None and caja['estado'] == 'CERRADA'
                    self.lbl_estado_text.config(
                        text="CERRADA HOY" if cerrada_hoy else "CERRADA",
                        fg=COLORES['gris_500'])
                    self.lbl_inicial.config(text="$ 0")
                    self.lbl_total_ventas.config(text="$ 0")
                    self.lbl_cuentas_abiertas.config(text="$ 0")
                    self.lbl_total_gastos.config(text="$ 0")
                    self.lbl_esperado.config(text="$ 0")

                    self.btn_apertura.config(state='normal')
                    self.btn_reporte_x.config(state='disabled')
                    self.btn_cierre.config(state='disabled')

                    for item in self.tree_movimientos.get_children():
                        self.tree_movimientos.delete(item)

        except Exception as e:
            messagebox.showerror("Error", f"Error al cargar caja: {str(e)}")

    def _abrir_caja(self):
        """Abre una caja nueva con diálogo propio (evita bugs de simpledialog)."""
        # Verificar si ya hay una caja abierta
        try:
            with conexion_segura() as conn:
                hoy = datetime.date.today().isoformat()
                existe = conn.execute(
                    "SELECT id_caja FROM caja_diaria WHERE estado='ABIERTA' AND DATE(fecha_apertura) = ?",
                    (hoy,)
                ).fetchone()
            if existe:
                messagebox.showwarning("Caja ya abierta", "Ya hay una caja abierta.")
                self._cargar_caja_actual()
                return
        except Exception as e:
            messagebox.showerror("Error", f"Error verificando caja: {e}")
            return

        # Diálogo propio con Entry
        dlg = tk.Toplevel(self.parent)
        dlg.title("Abrir Caja")
        dlg.resizable(False, False)
        dlg.grab_set()
        dlg.configure(bg=COLORES['fondo_card'])
        dlg.geometry("380x220")
        try:
            dlg.geometry(f"+{dlg.winfo_screenwidth()//2-190}+{dlg.winfo_screenheight()//2-110}")
        except Exception:
            pass

        tk.Frame(dlg, bg=COLORES['primario'], height=4).pack(fill='x')
        f = tk.Frame(dlg, bg=COLORES['fondo_card'], padx=24, pady=18)
        f.pack(fill='both', expand=True)

        tk.Label(f, text="Abrir Caja del Dia",
                 font=FUENTES['encabezado'], fg=COLORES['primario'],
                 bg=COLORES['fondo_card']).pack(anchor='w', pady=(0, 4))
        tk.Label(f, text="Ingrese el monto inicial en efectivo:",
                 font=FUENTES['normal'], fg=COLORES['texto'],
                 bg=COLORES['fondo_card']).pack(anchor='w', pady=(0, 6))

        entry = tk.Entry(f, font=FUENTES['precio'], relief='solid', bd=1,
                         bg=COLORES['fondo_input'], fg=COLORES['texto'],
                         insertbackground=COLORES['primario'])
        entry.pack(fill='x', ipady=8, pady=(0, 4))
        entry.insert(0, '0')
        entry.select_range(0, tk.END)
        entry.focus_set()

        lbl_err = tk.Label(f, text="", font=FUENTES['pequena'],
                           fg=COLORES['error'], bg=COLORES['fondo_card'])
        lbl_err.pack(anchor='w', pady=(0, 8))

        resultado = {'monto': None}

        def confirmar(event=None):
            txt = entry.get().strip().replace(',', '.').replace('$', '').replace(' ', '')
            try:
                monto = float(txt)
                if monto < 0:
                    raise ValueError
                resultado['monto'] = monto
                dlg.destroy()
            except ValueError:
                lbl_err.config(text="Ingrese un numero valido (>= 0)")

        def cancelar():
            dlg.destroy()

        entry.bind('<Return>', confirmar)
        btn_f = tk.Frame(f, bg=COLORES['fondo_card'])
        btn_f.pack(fill='x')
        crear_boton(btn_f, "Abrir Caja", confirmar, tipo='exito').pack(
            side='left', fill='x', expand=True, padx=(0, 6), ipady=4)
        crear_boton(btn_f, "Cancelar", cancelar, tipo='secundario').pack(
            side='left', fill='x', expand=True, ipady=4)

        dlg.wait_window()

        if resultado['monto'] is None:
            return

        monto = resultado['monto']
        try:
            with conexion_segura() as conn:
                conn.execute("""
                    INSERT INTO caja_diaria (
                        fecha_apertura, usuario_apertura, monto_inicial, estado, monto_esperado
                    ) VALUES (datetime('now','localtime'), ?, ?, 'ABIERTA', ?)
                """, (self.usuario['usuario'], monto, monto))

                id_nueva = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
                log_auditoria(conn, 'caja_diaria', id_nueva,
                              'INSERT', self.usuario['usuario'],
                              f"Apertura de caja: {format_money(monto)}")

            messagebox.showinfo("Caja Abierta", f"Caja abierta con {format_money(monto)}")
            self._cargar_caja_actual()

        except Exception as e:
            messagebox.showerror("Error", f"Error al abrir caja: {str(e)}")

    def _reporte_x(self):
        """Reporte X: muestra el resumen del turno actual SIN cerrar la caja."""
        if not self.caja_actual or self.caja_actual['estado'] != 'ABIERTA':
            messagebox.showwarning("Sin caja", "No hay caja abierta.")
            return

        hoy = datetime.date.today().isoformat()
        try:
            with conexion_segura() as conn:
                desglose = conn.execute("""
                    SELECT COALESCE(metodo_pago, 'Sin especificar') as metodo,
                           COUNT(*) as transacciones,
                           SUM(total) as total
                    FROM ventas
                    WHERE fecha_creacion >= ? AND fecha_creacion < date(?, '+1 day')
                      AND estado IN ('CERRADA', 'PAGADA')
                    GROUP BY metodo_pago
                    ORDER BY total DESC
                """, (hoy, hoy)).fetchall()

                total_ventas = conn.execute("""
                    SELECT COALESCE(SUM(total), 0) as t FROM ventas
                    WHERE fecha_creacion >= ? AND fecha_creacion < date(?, '+1 day')
                      AND estado IN ('CERRADA','PAGADA')
                """, (hoy, hoy)).fetchone()['t']

                total_gastos = conn.execute("""
                    SELECT COALESCE(SUM(valor), 0) as t FROM movimientos_caja
                    WHERE id_caja = ? AND tipo = 'EGRESO'
                """, (self.caja_actual['id_caja'],)).fetchone()['t']

                n_ventas = conn.execute("""
                    SELECT COUNT(*) as n FROM ventas
                    WHERE fecha_creacion >= ? AND fecha_creacion < date(?, '+1 day')
                      AND estado IN ('CERRADA','PAGADA')
                """, (hoy, hoy)).fetchone()['n']

        except Exception as e:
            messagebox.showerror("Error", f"Error al generar Reporte X: {e}")
            return

        # Diálogo informativo (solo lectura)
        dlg = tk.Toplevel(self.parent)
        dlg.title("Reporte X — Turno en Curso")
        dlg.geometry("480x460")
        dlg.transient(self.parent)
        dlg.grab_set()
        dlg.configure(bg=COLORES['fondo_card'])

        f = tk.Frame(dlg, bg=COLORES['fondo_card'], padx=22, pady=16)
        f.pack(fill='both', expand=True)

        # Encabezado
        hdr = tk.Frame(f, bg=COLORES['primario'])
        hdr.pack(fill='x', pady=(0, 12))
        tk.Label(hdr, text="REPORTE X — TURNO EN CURSO",
                 font=FUENTES['encabezado'], fg=COLORES['texto_claro'],
                 bg=COLORES['primario']).pack(pady=8)

        ahora = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        tk.Label(f, text=f"Generado: {ahora}",
                 font=FUENTES['normal'], fg=COLORES['texto_secundario'],
                 bg=COLORES['fondo_card']).pack(anchor='w')
        tk.Label(f, text=f"Apertura caja: {str(self.caja_actual['fecha_apertura'])[:16]}",
                 font=FUENTES['normal'], fg=COLORES['texto_secundario'],
                 bg=COLORES['fondo_card']).pack(anchor='w', pady=(0, 10))

        # KPIs rápidos
        kpi_f = tk.Frame(f, bg=COLORES['fondo_card'])
        kpi_f.pack(fill='x', pady=(0, 10))
        for lbl, val in [
            ("Monto inicial:", format_money(self.caja_actual['monto_inicial'])),
            ("Ventas del turno:", format_money(total_ventas)),
            (f"  ({n_ventas} transacciones)", ""),
            ("Gastos del turno:", f"-{format_money(total_gastos)}"),
            ("Monto esperado en caja:", format_money(self.caja_actual['monto_esperado'])),
        ]:
            row = tk.Frame(kpi_f, bg=COLORES['fondo_card'])
            row.pack(fill='x')
            tk.Label(row, text=lbl, font=FUENTES['normal'],
                     fg=COLORES['texto'], bg=COLORES['fondo_card'],
                     width=28, anchor='w').pack(side='left')
            if val:
                tk.Label(row, text=val, font=FUENTES['normal_bold'],
                         fg=COLORES['primario'], bg=COLORES['fondo_card']).pack(side='left')

        tk.Frame(f, bg=COLORES['borde'], height=1).pack(fill='x', pady=8)

        # Desglose por método
        tk.Label(f, text="Desglose por método de pago:",
                 font=FUENTES['normal_bold'], fg=COLORES['texto'],
                 bg=COLORES['fondo_card']).pack(anchor='w', pady=(0, 5))

        tree = ttk.Treeview(f, columns=('metodo', 'tx', 'total'),
                            show='headings', height=min(max(len(desglose), 1), 5))
        aplicar_estilo_tabla(tree)
        tree.heading('metodo', text='Método')
        tree.heading('tx', text='Tx')
        tree.heading('total', text='Total')
        tree.column('metodo', width=180, anchor='w')
        tree.column('tx', width=60, anchor='center')
        tree.column('total', width=130, anchor='e')
        for i, row in enumerate(desglose):
            tree.insert('', 'end', values=(row['metodo'], row['transacciones'],
                                           format_money(row['total'])),
                        tags=('par' if i % 2 == 0 else 'impar',))
        if not desglose:
            tree.insert('', 'end', values=('Sin ventas aun', '', ''))
        tree.pack(fill='x', pady=(0, 14))

        tk.Label(f, text="La caja SIGUE ABIERTA. Este reporte no cierra el turno.",
                 font=FUENTES['normal'], fg='#795548',
                 bg=COLORES['fondo_card']).pack(anchor='w', pady=(0, 8))

        crear_boton(f, "Cerrar", dlg.destroy, tipo='secundario').pack(anchor='e')

    def _cerrar_caja(self):
        """Cierra la caja del día (Reporte Z) mostrando desglose por método de pago."""
        if not self.caja_actual:
            messagebox.showwarning("Error", "No hay caja abierta")
            return

        hoy = datetime.date.today().isoformat()

        # ── Verificar cuentas abiertas pendientes ───────────────────
        try:
            with conexion_segura() as conn:
                cuentas_pendientes = conn.execute("""
                    SELECT COUNT(*) as n,
                           COALESCE(SUM(total), 0) as total_pendiente
                    FROM ventas
                    WHERE tipo = 'cuenta_abierta'
                      AND estado = 'ABIERTA'
                      AND DATE(fecha_creacion) = ?
                """, (hoy,)).fetchone()
        except Exception as e:
            messagebox.showerror("Error", f"Error al verificar cuentas: {str(e)}")
            return

        if cuentas_pendientes and cuentas_pendientes['n'] > 0:
            n = cuentas_pendientes['n']
            total_p = cuentas_pendientes['total_pendiente']
            messagebox.showerror(
                "No se puede cerrar la caja",
                f"Hay {n} cuenta(s) abierta(s) pendiente(s) por {format_money(total_p)}.\n\n"
                "Debe cobrar o cancelar todas las cuentas abiertas\n"
                "antes de cerrar la caja del día."
            )
            return

        # ── Obtener desglose de ventas del día por método de pago ───
        try:
            with conexion_segura() as conn:
                desglose = conn.execute("""
                    SELECT COALESCE(metodo_pago, 'Sin especificar') as metodo,
                           COUNT(*) as transacciones,
                           SUM(total) as total
                    FROM ventas
                    WHERE DATE(fecha_creacion) = ?
                      AND estado IN ('CERRADA', 'PAGADA')
                    GROUP BY metodo_pago
                    ORDER BY total DESC
                """, (hoy,)).fetchall()
        except Exception as e:
            messagebox.showerror("Error", f"Error al obtener desglose: {str(e)}")
            return

        # ── Calcular totales incluyendo gastos y desglose por método ──
        try:
            with conexion_segura() as conn:
                total_ventas_hoy = conn.execute("""
                    SELECT COALESCE(SUM(total), 0) as t FROM ventas
                    WHERE DATE(fecha_creacion) = ?
                      AND estado IN ('CERRADA', 'PAGADA')
                """, (hoy,)).fetchone()['t']

                # Gastos reales del día — excluye devoluciones Y anulaciones
                # (las anulaciones ya están excluidas del total_ventas_hoy;
                #  incluirlas aquí también las descontaría dos veces)
                total_gastos_real = conn.execute("""
                    SELECT COALESCE(SUM(valor), 0) as t FROM movimientos_caja
                    WHERE id_caja = ? AND tipo = 'EGRESO'
                      AND concepto NOT LIKE 'Devolucion venta%'
                      AND concepto NOT LIKE 'Anulacion venta%'
                """, (self.caja_actual['id_caja'],)).fetchone()['t']

                # Total anulaciones del día (informativo, no afecta monto_esperado)
                total_anulaciones = conn.execute("""
                    SELECT COALESCE(SUM(total), 0) FROM ventas
                    WHERE DATE(fecha_creacion) = ? AND estado = 'ANULADA'
                """, (hoy,)).fetchone()[0]

                # Solo devoluciones que sí salieron como efectivo de esta caja
                # (aplicado_caja=0 → la caja estaba cerrada cuando se registró, no afecta el saldo)
                total_devoluciones = conn.execute("""
                    SELECT COALESCE(SUM(monto), 0) FROM devoluciones
                    WHERE DATE(fecha) = ? AND aplicado_caja = 1
                """, (hoy,)).fetchone()[0]

                # Efectivo recibido hoy (incluye efectivo de ventas MIXTO)
                efectivo_ventas = conn.execute("""
                    SELECT COALESCE(SUM(p.valor), 0) FROM pagos p
                    JOIN ventas v ON p.id_venta = v.id_venta
                    WHERE DATE(v.fecha_creacion) = ?
                      AND v.estado IN ('CERRADA','PAGADA')
                      AND p.metodo_pago = 'EFECTIVO'
                """, (hoy,)).fetchone()[0]

                # Transferencias recibidas hoy (Nequi + TuLlave)
                transfer_ventas = conn.execute("""
                    SELECT COALESCE(SUM(p.valor), 0) FROM pagos p
                    JOIN ventas v ON p.id_venta = v.id_venta
                    WHERE DATE(v.fecha_creacion) = ?
                      AND v.estado IN ('CERRADA','PAGADA')
                      AND p.metodo_pago IN ('NEQUI','TULLAVE')
                """, (hoy,)).fetchone()[0]

                monto_esperado_real = (
                    self.caja_actual['monto_inicial'] + total_ventas_hoy - total_gastos_real - total_devoluciones
                )
                # Efectivo esperado SIN la base = efectivo recibido menos gastos y devoluciones
                efectivo_sin_base = efectivo_ventas - total_gastos_real - total_devoluciones
        except Exception as e:
            _log.error(f"Error calculando totales para Reporte Z: {e}")
            monto_esperado_real = self.caja_actual['monto_esperado']
            total_gastos_real = 0
            total_anulaciones = 0
            efectivo_ventas = 0
            transfer_ventas = 0
            efectivo_sin_base = 0

        # Construir diálogo con desglose + input de monto contado
        dialogo = tk.Toplevel(self.parent)
        dialogo.title("Reporte Z — Cierre Definitivo del Dia")
        dialogo.geometry("520x600")
        dialogo.transient(self.parent)
        dialogo.grab_set()
        dialogo.configure(bg=COLORES['fondo_card'])

        main_frame = tk.Frame(dialogo, bg=COLORES['fondo_card'], padx=20, pady=15)
        main_frame.pack(fill='both', expand=True)

        # Título
        hdr = tk.Frame(main_frame, bg=COLORES['primario'])
        hdr.pack(fill='x', pady=(0, 12))
        tk.Label(hdr, text="REPORTE Z — CIERRE DEFINITIVO DEL DIA",
                 font=FUENTES['encabezado'], fg=COLORES['texto_claro'],
                 bg=COLORES['primario']).pack(pady=8)

        # KPIs principales
        kpi_f = tk.Frame(main_frame, bg=COLORES['fondo_card'])
        kpi_f.pack(fill='x', pady=(0, 8))
        def _fila_kpi(parent, lbl, val, color=None):
            row = tk.Frame(parent, bg=COLORES['fondo_card'])
            row.pack(fill='x', pady=2)
            tk.Label(row, text=lbl, font=FUENTES['normal'],
                     fg=COLORES['texto'], bg=COLORES['fondo_card'],
                     width=30, anchor='w').pack(side='left')
            tk.Label(row, text=val, font=FUENTES['normal_bold'],
                     fg=color or COLORES['primario'],
                     bg=COLORES['fondo_card']).pack(side='left')

        _fila_kpi(kpi_f, "Total vendido hoy:",           format_money(total_ventas_hoy))
        _fila_kpi(kpi_f, "  Cobrado en efectivo:",       format_money(efectivo_ventas))
        _fila_kpi(kpi_f, "  Cobrado por transferencia:", format_money(transfer_ventas))
        if total_gastos_real > 0:
            _fila_kpi(kpi_f, "Gastos del dia:",          f"-{format_money(total_gastos_real)}",
                      COLORES['error'])
        if total_devoluciones > 0:
            _fila_kpi(kpi_f, "Devoluciones del dia:",    f"-{format_money(total_devoluciones)}",
                      COLORES['error'])
        if total_anulaciones > 0:
            _fila_kpi(kpi_f, "Ventas anuladas (info):",  format_money(total_anulaciones),
                      COLORES['texto_secundario'])

        tk.Frame(main_frame, bg=COLORES['acento'], height=2).pack(fill='x', pady=6)

        # Cuadro destacado: efectivo a contar SIN base
        ef_frame = tk.Frame(main_frame, bg='#e8f5e9', relief='solid', bd=1)
        ef_frame.pack(fill='x', pady=(0, 8))
        tk.Label(ef_frame, text="EFECTIVO A CONTAR (sin contar la base):",
                 font=FUENTES['normal_bold'], fg=COLORES['primario'],
                 bg='#e8f5e9').pack(anchor='w', padx=10, pady=(6, 0))
        tk.Label(ef_frame, text=format_money(max(efectivo_sin_base, 0)),
                 font=('Consolas', 20, 'bold'), fg=COLORES['primario'],
                 bg='#e8f5e9').pack(anchor='w', padx=10, pady=(0, 6))
        tk.Label(ef_frame,
                 text=f"(Base inicial {format_money(self.caja_actual['monto_inicial'])} NO incluida)",
                 font=FUENTES['pequena'], fg=COLORES['texto_secundario'],
                 bg='#e8f5e9').pack(anchor='w', padx=10, pady=(0, 6))

        # Tabla de desglose por método de pago
        tk.Label(main_frame, text="Desglose por método de pago:",
                 font=FUENTES['normal_bold'], fg=COLORES['texto'],
                 bg=COLORES['fondo_card']).pack(anchor='w', pady=(0, 5))

        tree_frame = tk.Frame(main_frame, bg=COLORES['fondo_card'])
        tree_frame.pack(fill='x', pady=(0, 12))

        tree = ttk.Treeview(tree_frame,
                            columns=('metodo', 'transacciones', 'total'),
                            show='headings',
                            height=min(max(len(desglose), 1), 6))
        tree.column('metodo', width=180, anchor='w')
        tree.column('transacciones', width=110, anchor='center')
        tree.column('total', width=150, anchor='e')
        tree.heading('metodo', text='Método de Pago')
        tree.heading('transacciones', text='Transacciones')
        tree.heading('total', text='Total')
        aplicar_estilo_tabla(tree)

        if desglose:
            for i, row in enumerate(desglose):
                tag = 'par' if i % 2 == 0 else 'impar'
                tree.insert('', 'end', values=(
                    row['metodo'], row['transacciones'], format_money(row['total'])
                ), tags=(tag,))
        else:
            tree.insert('', 'end', values=('Sin ventas registradas hoy', '', ''))

        tree.pack(fill='x')

        # Separador
        tk.Frame(main_frame, bg=COLORES['borde'], height=1).pack(fill='x', pady=10)

        # Input de efectivo contado
        tk.Label(main_frame, text=" Ingrese el efectivo contado en caja:",
                 font=FUENTES['normal_bold'], fg=COLORES['texto'],
                 bg=COLORES['fondo_card']).pack(anchor='w', pady=(0, 5))

        entry_monto = tk.Entry(main_frame, font=FUENTES['precio'],
                               relief='solid', bd=1, highlightthickness=1,
                               highlightbackground=COLORES['borde'],
                               bg=COLORES['fondo_input'], fg=COLORES['texto'],
                               insertbackground=COLORES['acento'])
        entry_monto.pack(fill='x', ipady=8, pady=(0, 15))
        entry_monto.focus_set()

        resultado = {'monto': None}

        def confirmar(event=None):
            monto_str = entry_monto.get().strip().replace(',', '.')
            try:
                monto = float(monto_str)
                if monto < 0:
                    raise ValueError("Negativo")
                resultado['monto'] = monto
                dialogo.destroy()
            except ValueError:
                messagebox.showerror("Error", "Ingrese un número válido (>= 0)", parent=dialogo)

        def cancelar():
            dialogo.destroy()

        entry_monto.bind('<Return>', confirmar)

        btn_frame = tk.Frame(main_frame, bg=COLORES['fondo_card'])
        btn_frame.pack(fill='x')
        crear_boton(btn_frame, " Confirmar Cierre", confirmar, tipo='exito').pack(
            side='left', padx=(0, 6), fill='x', expand=True)
        crear_boton(btn_frame, " Cancelar", cancelar, tipo='secundario').pack(
            side='left', fill='x', expand=True)

        dialogo.wait_window()

        if resultado['monto'] is None:
            return

        monto_real = resultado['monto']

        try:
            with conexion_segura() as conn:
                diferencia = monto_real - monto_esperado_real

                conn.execute("""
                    UPDATE caja_diaria
                    SET fecha_cierre = datetime('now','localtime'),
                        usuario_cierre = ?,
                        monto_real = ?,
                        diferencia = ?,
                        monto_esperado = ?,
                        estado = 'CERRADA'
                    WHERE id_caja = ?
                """, (self.usuario['usuario'], monto_real, diferencia,
                      monto_esperado_real, self.caja_actual['id_caja']))

                log_auditoria(conn, 'caja_diaria', self.caja_actual['id_caja'], 'UPDATE',
                              self.usuario['usuario'],
                              f"Cierre de caja. Diferencia: {format_money(diferencia)}")

            msg = (f"Caja cerrada\n\n"
                   f"Esperado:  {format_money(monto_esperado_real)}\n"
                   f"Contado:   {format_money(monto_real)}\n\n")
            if diferencia > 0:
                msg += f" Sobrante: +{format_money(diferencia)}"
            elif diferencia < 0:
                msg += f"AVISO: Faltante: {format_money(abs(diferencia))}"
            else:
                msg += " Cuadra perfecto"

            messagebox.showinfo(" Caja Cerrada", msg)
            self._cargar_caja_actual()

            # Ofrecer imprimir resumen del día
            if messagebox.askyesno(
                "Resumen del dia",
                "¿Desea ver e imprimir el resumen de ventas del dia?"
            ):
                self._mostrar_resumen_dia(hoy)

            # Backup automático al cerrar caja (Reporte Z)
            try:
                import threading as _thr
                import shutil as _shutil
                import os as _os
                from database.connection import get_db_path, conexion_segura as _cs

                def _backup_z():
                    try:
                        with _cs() as conn_b:
                            ruta_val = conn_b.execute(
                                "SELECT valor FROM configuracion WHERE clave = 'backup_ruta'"
                            ).fetchone()
                        ruta = ruta_val['valor'].strip() if ruta_val and ruta_val['valor'] else ''
                        if not ruta or not _os.path.isdir(ruta):
                            return
                        import datetime as _dt
                        ts = _dt.datetime.now().strftime('%Y%m%d_%H%M%S')
                        dest = _os.path.join(ruta, f"backup_cierre_caja_{ts}.db")
                        import sqlite3 as _sq
                        src_conn = _sq.connect(get_db_path())
                        dst_conn = _sq.connect(dest)
                        src_conn.backup(dst_conn)
                        src_conn.close()
                        dst_conn.close()
                    except Exception as e:
                        _log.error(f"Error creando backup al cerrar caja: {e}")

                _thr.Thread(target=_backup_z, daemon=True, name="BackupCierreCaja").start()
            except Exception as e:
                _log.error(f"Error iniciando thread de backup cierre caja: {e}")

        except Exception as e:
            messagebox.showerror("Error", f"Error al cerrar caja: {str(e)}")

    def _mostrar_resumen_dia(self, fecha):
        """Muestra ventana con todas las ventas del día y permite exportar/imprimir."""
        try:
            with conexion_segura() as conn:
                ventas = conn.execute("""
                    SELECT v.numero_venta, v.fecha_creacion, v.tipo, v.metodo_pago,
                           v.total, v.estado,
                           COALESCE(u.nombre_completo, v.usuario_nombre) as cajero
                    FROM ventas v
                    LEFT JOIN usuarios u ON u.id_usuario = v.id_usuario
                    WHERE DATE(v.fecha_creacion) = ?
                      AND v.estado != 'ANULADA'
                    ORDER BY v.fecha_creacion
                """, (fecha,)).fetchall()

                anuladas = conn.execute("""
                    SELECT numero_venta, total, fecha_creacion,
                           COALESCE(usuario_nombre, '') as cajero
                    FROM ventas
                    WHERE DATE(fecha_creacion) = ? AND estado = 'ANULADA'
                    ORDER BY fecha_creacion
                """, (fecha,)).fetchall()

                devoluciones = conn.execute("""
                    SELECT d.monto, d.motivo, d.usuario, d.fecha,
                           v.numero_venta
                    FROM devoluciones d
                    JOIN ventas v ON v.id_venta = d.id_venta
                    WHERE DATE(d.fecha) = ?
                    ORDER BY d.fecha
                """, (fecha,)).fetchall()

                gastos = conn.execute("""
                    SELECT descripcion, valor, categoria, usuario_registro, fecha
                    FROM gastos
                    WHERE DATE(fecha) = ?
                    ORDER BY fecha
                """, (fecha,)).fetchall()

                boletas = conn.execute("""
                    SELECT numero_boleta, hora_entrada, cantidad_personas, total
                    FROM boletas_entrada
                    WHERE DATE(hora_entrada) = ?
                    ORDER BY hora_entrada
                """, (fecha,)).fetchall()
        except Exception as e:
            messagebox.showerror("Error", f"Error al cargar resumen: {e}")
            return

        # Consumo de empleados (tabla puede no existir en instalaciones viejas)
        consumos = []
        try:
            with conexion_segura() as conn:
                consumos = conn.execute("""
                    SELECT c.total, c.fecha,
                           COALESCE(e.nombre, 'Empleado') AS empleado,
                           (SELECT GROUP_CONCAT(d.producto_nombre || ' x' ||
                                CAST(CAST(d.cantidad AS INTEGER) AS TEXT), ', ')
                            FROM consumo_detalle d WHERE d.id_consumo = c.id_consumo
                           ) AS items
                    FROM consumos_empleados c
                    LEFT JOIN nomina_empleados e ON e.id_empleado = c.id_empleado
                    WHERE DATE(c.fecha) = ? AND c.estado = 'CERRADO'
                    ORDER BY c.fecha
                """, (fecha,)).fetchall()
        except Exception:
            pass  # tabla aun no existe en esta BD

        total_v   = sum(r['total'] for r in ventas)
        total_b   = sum(r['total'] for r in boletas)
        total_g   = sum(r['valor'] for r in gastos)
        total_dev = sum(r['monto'] for r in devoluciones)
        total_anu = sum(r['total'] for r in anuladas)
        total_con = sum(r['total'] for r in consumos)
        neto = (total_v - total_dev) + total_b - total_g

        dlg = tk.Toplevel(self.parent)
        dlg.title(f"Resumen del dia — {fecha}")
        dlg.configure(bg=COLORES['fondo'])
        dlg.geometry("820x680")
        dlg.transient(self.parent)
        dlg.grab_set()

        # Header
        hdr = tk.Frame(dlg, bg=COLORES['primario'], pady=10)
        hdr.pack(fill='x')
        tk.Label(hdr, text=f"RESUMEN DEL DIA — {fecha}",
                 font=FUENTES['encabezado'], fg=COLORES['texto_claro'],
                 bg=COLORES['primario']).pack()
        tk.Label(hdr, text="Club Los Pocitos Azufrados",
                 font=FUENTES['normal'], fg=COLORES['acento'],
                 bg=COLORES['primario']).pack()

        # Totales rápidos
        kpi_frame = tk.Frame(dlg, bg=COLORES['fondo'], pady=8)
        kpi_frame.pack(fill='x', padx=15)
        kpi_items = [
            ("Ventas netas",   format_money(total_v - total_dev), COLORES['exito']),
            ("Boletas",        format_money(total_b),             COLORES['agua']),
            ("Devoluciones",   format_money(total_dev),           COLORES['advertencia']),
            ("Anuladas",       format_money(total_anu),           COLORES['error']),
            ("Gastos",         format_money(total_g),             COLORES['error']),
            ("Consumo emp.",   format_money(total_con),           COLORES['texto_secundario']),
            ("NETO",           format_money(neto),                COLORES['primario']),
        ]
        for titulo, valor, color in kpi_items:
            card = tk.Frame(kpi_frame, bg=COLORES['fondo_card'],
                            relief='solid', bd=1, padx=10, pady=6)
            card.pack(side='left', expand=True, fill='both', padx=4)
            tk.Frame(card, bg=color, height=3).pack(fill='x')
            tk.Label(card, text=titulo, font=FUENTES['pequena'],
                     fg=COLORES['texto_secundario'], bg=COLORES['fondo_card']).pack()
            tk.Label(card, text=valor, font=FUENTES['normal_bold'],
                     fg=color, bg=COLORES['fondo_card']).pack()

        # Notebook con tabs
        nb = ttk.Notebook(dlg)
        nb.pack(fill='both', expand=True, padx=15, pady=10)

        # Tab ventas
        tab_v = tk.Frame(nb, bg=COLORES['fondo'])
        nb.add(tab_v, text=f"Ventas ({len(ventas)})")
        tree_v = ttk.Treeview(tab_v,
                               columns=('num', 'hora', 'tipo', 'metodo', 'cajero', 'total'),
                               show='headings')
        tree_v.heading('num',    text='N. Venta')
        tree_v.heading('hora',   text='Hora')
        tree_v.heading('tipo',   text='Tipo')
        tree_v.heading('metodo', text='Pago')
        tree_v.heading('cajero', text='Cajero')
        tree_v.heading('total',  text='Total')
        tree_v.column('num',    width=120)
        tree_v.column('hora',   width=80)
        tree_v.column('tipo',   width=100)
        tree_v.column('metodo', width=100)
        tree_v.column('cajero', width=130)
        tree_v.column('total',  width=110, anchor='e')
        aplicar_estilo_tabla(tree_v)
        sb_v = ttk.Scrollbar(tab_v, orient='vertical', command=tree_v.yview)
        tree_v.configure(yscrollcommand=sb_v.set)
        sb_v.pack(side='right', fill='y')
        tree_v.pack(fill='both', expand=True)
        for i, r in enumerate(ventas):
            hora = r['fecha_creacion'][11:16] if r['fecha_creacion'] else ''
            tag = 'par' if i % 2 == 0 else 'impar'
            tree_v.insert('', 'end', values=(
                r['numero_venta'], hora, r['tipo'], r['metodo_pago'] or '-',
                r['cajero'] or '-', format_money(r['total'])
            ), tags=(tag,))

        # Tab boletas
        tab_b = tk.Frame(nb, bg=COLORES['fondo'])
        nb.add(tab_b, text=f"Boletas ({len(boletas)})")
        tree_b = ttk.Treeview(tab_b,
                               columns=('num', 'hora', 'personas', 'total'),
                               show='headings')
        tree_b.heading('num',      text='N. Boleta')
        tree_b.heading('hora',     text='Hora')
        tree_b.heading('personas', text='Personas')
        tree_b.heading('total',    text='Total')
        tree_b.column('num',      width=150)
        tree_b.column('hora',     width=100)
        tree_b.column('personas', width=100, anchor='center')
        tree_b.column('total',    width=110, anchor='e')
        aplicar_estilo_tabla(tree_b)
        tree_b.pack(fill='both', expand=True)
        for i, r in enumerate(boletas):
            hora = r['hora_entrada'][11:16] if r['hora_entrada'] else ''
            tag = 'par' if i % 2 == 0 else 'impar'
            tree_b.insert('', 'end', values=(
                r['numero_boleta'], hora, r['cantidad_personas'], format_money(r['total'])
            ), tags=(tag,))

        # Tab gastos
        tab_g = tk.Frame(nb, bg=COLORES['fondo'])
        nb.add(tab_g, text=f"Gastos ({len(gastos)})")
        tree_g = ttk.Treeview(tab_g,
                               columns=('desc', 'categoria', 'usuario', 'valor'),
                               show='headings')
        tree_g.heading('desc',      text='Descripcion')
        tree_g.heading('categoria', text='Categoria')
        tree_g.heading('usuario',   text='Registrado por')
        tree_g.heading('valor',     text='Valor')
        tree_g.column('desc',      width=220)
        tree_g.column('categoria', width=120)
        tree_g.column('usuario',   width=120)
        tree_g.column('valor',     width=110, anchor='e')
        aplicar_estilo_tabla(tree_g)
        tree_g.pack(fill='both', expand=True)
        for i, r in enumerate(gastos):
            tag = 'par' if i % 2 == 0 else 'impar'
            tree_g.insert('', 'end', values=(
                r['descripcion'], r['categoria'] or '-', r['usuario_registro'] or '-',
                format_money(r['valor'])
            ), tags=(tag,))

        # Tab devoluciones + anulaciones
        tab_dev = tk.Frame(nb, bg=COLORES['fondo'])
        nb.add(tab_dev, text=f"Devoluciones/Anuladas ({len(devoluciones) + len(anuladas)})")
        tree_dev = ttk.Treeview(tab_dev,
                                columns=('tipo', 'referencia', 'motivo', 'usuario', 'monto'),
                                show='headings')
        tree_dev.heading('tipo',       text='Tipo')
        tree_dev.heading('referencia', text='N. Venta')
        tree_dev.heading('motivo',     text='Motivo')
        tree_dev.heading('usuario',    text='Usuario')
        tree_dev.heading('monto',      text='Monto')
        tree_dev.column('tipo',       width=90,  anchor='center')
        tree_dev.column('referencia', width=130)
        tree_dev.column('motivo',     width=220)
        tree_dev.column('usuario',    width=110)
        tree_dev.column('monto',      width=100, anchor='e')
        aplicar_estilo_tabla(tree_dev)
        tree_dev.pack(fill='both', expand=True)
        for i, r in enumerate(devoluciones):
            tag = 'par' if i % 2 == 0 else 'impar'
            tree_dev.insert('', 'end', tags=(tag,), values=(
                'Devolucion', r['numero_venta'],
                (r['motivo'] or '')[:40], r['usuario'] or '-',
                format_money(r['monto'])))
        for i, r in enumerate(anuladas):
            tag = 'par' if (i + len(devoluciones)) % 2 == 0 else 'impar'
            tree_dev.insert('', 'end', tags=(tag,), values=(
                'ANULADA', r['numero_venta'],
                'Venta anulada', r['cajero'] or '-',
                format_money(r['total'])))

        # Tab consumo empleados
        tab_con = tk.Frame(nb, bg=COLORES['fondo'])
        nb.add(tab_con, text=f"Consumo Empleados ({len(consumos)})")
        tree_con = ttk.Treeview(tab_con,
                                columns=('empleado', 'hora', 'items', 'total'),
                                show='headings')
        tree_con.heading('empleado', text='Empleado')
        tree_con.heading('hora',     text='Hora')
        tree_con.heading('items',    text='Productos')
        tree_con.heading('total',    text='Total (sin cobro)')
        tree_con.column('empleado', width=140)
        tree_con.column('hora',     width=70, anchor='center')
        tree_con.column('items',    width=300)
        tree_con.column('total',    width=120, anchor='e')
        aplicar_estilo_tabla(tree_con)
        tree_con.pack(fill='both', expand=True)
        for i, r in enumerate(consumos):
            hora = r['fecha'][11:16] if r['fecha'] else ''
            tag = 'par' if i % 2 == 0 else 'impar'
            tree_con.insert('', 'end', tags=(tag,), values=(
                r['empleado'], hora,
                r['items'] or '-', format_money(r['total'])))
        if not consumos:
            tk.Label(tab_con, text="Sin consumo de empleados registrado hoy.",
                     font=FUENTES['normal'], fg=COLORES['texto_secundario'],
                     bg=COLORES['fondo']).pack(pady=20)

        # Botones
        btn_frame = tk.Frame(dlg, bg=COLORES['fondo'], pady=8)
        btn_frame.pack(fill='x', padx=15)

        def _exportar_excel():
            try:
                import openpyxl
                from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            except ImportError:
                messagebox.showerror("Error", "Instale openpyxl: pip install openpyxl", parent=dlg)
                return
            wb = openpyxl.Workbook()
            hf = Font(bold=True, color="FFFFFF")
            hfill = PatternFill("solid", fgColor="1B5E20")
            afill = PatternFill("solid", fgColor="E8F5E9")

            def _header(ws, row, cols):
                for c in range(1, cols + 1):
                    ws.cell(row=row, column=c).font = hf
                    ws.cell(row=row, column=c).fill = hfill

            # Hoja resumen
            ws = wb.active
            ws.title = "Resumen"
            ws["A1"] = f"Resumen del dia — {fecha}"
            ws["A1"].font = Font(bold=True, size=13, color="1B5E20")
            ws.append([])
            ws.append(["Concepto", "Valor"])
            _header(ws, 3, 2)
            for concepto, valor in [
                ("Total Ventas", total_v), ("Total Boletas", total_b),
                ("Total Gastos", total_g), ("NETO del dia", neto)
            ]:
                ws.append([concepto, valor])

            # Hoja ventas
            ws2 = wb.create_sheet("Ventas")
            ws2.append(["N. Venta", "Hora", "Tipo", "Metodo Pago", "Cajero", "Total"])
            _header(ws2, 1, 6)
            for i, r in enumerate(ventas):
                hora = r['fecha_creacion'][11:16] if r['fecha_creacion'] else ''
                ws2.append([r['numero_venta'], hora, r['tipo'],
                            r['metodo_pago'] or '-', r['cajero'] or '-', r['total']])
                if i % 2 == 0:
                    for c in range(1, 7):
                        ws2.cell(row=i + 2, column=c).fill = afill

            # Hoja gastos
            ws3 = wb.create_sheet("Gastos")
            ws3.append(["Descripcion", "Categoria", "Usuario", "Valor"])
            _header(ws3, 1, 4)
            for i, r in enumerate(gastos):
                ws3.append([r['descripcion'], r['categoria'] or '-',
                            r['usuario_registro'] or '-', r['valor']])
                if i % 2 == 0:
                    for c in range(1, 5):
                        ws3.cell(row=i + 2, column=c).fill = afill

            for ws_s in [ws, ws2, ws3]:
                for col in ws_s.columns:
                    ws_s.column_dimensions[col[0].column_letter].width = min(
                        max((len(str(c.value or '')) for c in col), default=10) + 4, 50)

            nombre = f"resumen_dia_{fecha}.xlsx"
            ruta = os.path.join(BASE_DIR, 'data', nombre)
            wb.save(ruta)
            messagebox.showinfo("Exportado", f"Guardado en:\n{ruta}", parent=dlg)
            try:
                import subprocess
                subprocess.Popen(['start', '', ruta], shell=True)
            except Exception:
                pass

        crear_boton(btn_frame, "Exportar Excel", _exportar_excel, tipo='agua').pack(
            side='left', padx=(0, 8), ipady=4)
        crear_boton(btn_frame, "Cerrar", dlg.destroy, tipo='secundario').pack(
            side='left', ipady=4)

    def _ir_al_bar(self):
        if self._callback_ir_bar:
            self._callback_ir_bar()

    def _registrar_gasto_rapido(self):
        """Abre un diálogo para registrar un gasto rápido"""
        if not self.caja_actual or self.caja_actual['estado'] != 'ABIERTA':
            messagebox.showwarning("Error", "Debe tener una caja abierta para registrar gastos")
            return

        # Crear diálogo para gasto rápido
        dialogo = tk.Toplevel(self.parent)
        dialogo.title("Registrar Gasto Rápido")
        dialogo.geometry("400x300")
        dialogo.transient(self.parent)

        # Frame principal
        main_frame = tk.Frame(dialogo, bg=COLORES['fondo_card'])
        main_frame.pack(fill='both', expand=True, padx=15, pady=15)

        # Descripción
        tk.Label(main_frame, text="Descripción:", font=FUENTES['normal'],
                 fg=COLORES['texto'], bg=COLORES['fondo_card']).pack(anchor='w', pady=(0, 5))
        entry_desc = tk.Entry(main_frame, font=FUENTES['input'],
                             relief='solid', bd=1)
        entry_desc.pack(fill='x', ipady=6, pady=(0, 10))

        # Monto
        tk.Label(main_frame, text="Monto ($):", font=FUENTES['normal'],
                 fg=COLORES['texto'], bg=COLORES['fondo_card']).pack(anchor='w', pady=(0, 5))
        entry_monto = tk.Entry(main_frame, font=FUENTES['input'],
                              relief='solid', bd=1)
        entry_monto.pack(fill='x', ipady=6, pady=(0, 10))

        # Categoría
        tk.Label(main_frame, text="Categoría:", font=FUENTES['normal'],
                 fg=COLORES['texto'], bg=COLORES['fondo_card']).pack(anchor='w', pady=(0, 5))
        combo_cat = ttk.Combobox(main_frame, values=['Otros', 'Compras Insumos', 'Servicios Públicos', 'Mantenimiento', 'Transporte'],
                                state='readonly', font=FUENTES['input'])
        combo_cat.set('Otros')
        combo_cat.pack(fill='x', ipady=6, pady=(0, 15))

        # Botones
        btn_frame = tk.Frame(main_frame, bg=COLORES['fondo_card'])
        btn_frame.pack(fill='x')

        def guardar_gasto():
            descripcion = entry_desc.get().strip()
            monto_str = entry_monto.get().strip()

            if not descripcion or not monto_str:
                messagebox.showwarning("Campos Requeridos", "Complete descripción y monto")
                return

            try:
                monto = float(monto_str)
                if monto <= 0:
                    raise ValueError()
            except ValueError:
                messagebox.showerror("Error", "Ingrese un monto válido")
                return

            try:
                with conexion_segura() as conn:
                    # Insertar gasto
                    conn.execute("""
                        INSERT INTO gastos (
                            fecha, descripcion, valor, categoria, metodo_pago,
                            usuario_registro, tipo_gasto
                        ) VALUES (datetime('now','localtime'), ?, ?, ?, 'EFECTIVO', ?, 'OPERATIVO')
                    """, (descripcion, monto, combo_cat.get(), self.usuario['usuario']))

                    gasto_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]

                    # Actualizar movimiento de caja
                    conn.execute("""
                        INSERT INTO movimientos_caja (
                            id_caja, tipo, concepto, valor, usuario
                        ) VALUES (?, 'EGRESO', ?, ?, ?)
                    """, (self.caja_actual['id_caja'], descripcion, monto, self.usuario['usuario']))

                    # Actualizar total_gastos en caja_diaria
                    nuevo_total = (self.caja_actual['total_gastos'] or 0) + monto
                    conn.execute("""
                        UPDATE caja_diaria
                        SET total_gastos = ?, monto_esperado = monto_inicial + total_ventas - ?
                        WHERE id_caja = ?
                    """, (nuevo_total, nuevo_total, self.caja_actual['id_caja']))

                    log_auditoria(conn, 'gastos', gasto_id, 'INSERT',
                                self.usuario['usuario'], f"Gasto rápido: {descripcion} - {format_money(monto)}")

                messagebox.showinfo("Exito", f"Gasto registrado: {format_money(monto)}")
                dialogo.destroy()
                self._cargar_caja_actual()

            except Exception as e:
                messagebox.showerror("Error", f"Error al guardar gasto: {str(e)}")

        crear_boton(btn_frame, " Guardar", guardar_gasto, tipo='exito').pack(side='left', padx=(0, 5), fill='x', expand=True)
        crear_boton(btn_frame, " Cancelar", dialogo.destroy, tipo='secundario').pack(side='left', fill='x', expand=True)

    def detener(self):
        try:
            self._canvas_caja.unbind_all('<MouseWheel>')
            self._canvas_caja.unbind_all('<Button-4>')
            self._canvas_caja.unbind_all('<Button-5>')
        except Exception:
            pass
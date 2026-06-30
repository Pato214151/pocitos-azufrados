"""
Módulo de Reportes - Club Los Pocitos Azufrados
Reportes de ventas, ingresos vs gastos y estadísticas
"""

import tkinter as tk
from tkinter import ttk, messagebox
import os, sys, datetime

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE_DIR)

from utils.tema_corporativo import COLORES, FUENTES, crear_boton, format_money, aplicar_estilo_tabla
from database.connection import get_connection, conexion_segura


class ReportesModule:
    def __init__(self, parent, usuario):
        self.parent = parent
        self.usuario = usuario
        self.fecha_inicio = datetime.date.today()
        self.fecha_fin = datetime.date.today()

        self._crear_interfaz()
        self._generar_reportes()

    def _crear_interfaz(self):
        """Interfaz principal de reportes"""
        # Contenedor de cabecera
        header_frame = tk.Frame(self.parent, bg=COLORES['fondo'])
        header_frame.pack(fill='x', padx=15, pady=(10, 5))

        # Fila 1: Título + entradas de fecha
        header = tk.Frame(header_frame, bg=COLORES['fondo'])
        header.pack(fill='x')

        tk.Label(header, text="Reportes",
                 font=FUENTES['titulo'], fg=COLORES['texto'],
                 bg=COLORES['fondo']).pack(side='left')

        filtro_frame = tk.Frame(header, bg=COLORES['fondo'])
        filtro_frame.pack(side='right')

        tk.Label(filtro_frame, text="Desde:", font=FUENTES['pequena'],
                 fg=COLORES['texto_secundario'], bg=COLORES['fondo']).pack(side='left', padx=5)
        self.entry_desde = tk.Entry(filtro_frame, font=FUENTES['input'], width=12)
        self.entry_desde.insert(0, self.fecha_inicio.isoformat())
        self.entry_desde.pack(side='left', padx=5)

        tk.Label(filtro_frame, text="Hasta:", font=FUENTES['pequena'],
                 fg=COLORES['texto_secundario'], bg=COLORES['fondo']).pack(side='left', padx=5)
        self.entry_hasta = tk.Entry(filtro_frame, font=FUENTES['input'], width=12)
        self.entry_hasta.insert(0, self.fecha_fin.isoformat())
        self.entry_hasta.pack(side='left', padx=5)

        crear_boton(filtro_frame, "Actualizar", self._generar_reportes, tipo='primario').pack(side='left', padx=5)
        crear_boton(filtro_frame, "Exportar Excel", self._exportar_excel, tipo='agua').pack(side='left', padx=5)

        # Fila 2: Botones rápidos de periodo
        rapidos_frame = tk.Frame(header_frame, bg=COLORES['fondo'])
        rapidos_frame.pack(fill='x', pady=(6, 0))

        tk.Label(rapidos_frame, text="Periodo:", font=FUENTES['pequena'],
                 fg=COLORES['texto_deshabilitado'], bg=COLORES['fondo']).pack(side='left', padx=(0, 8))

        periodos = [
            ("Hoy",          'hoy'),
            ("Esta semana",  'semana'),
            ("Este mes",     'mes'),
            ("Mes anterior", 'mes_anterior'),
        ]
        for label, clave in periodos:
            crear_boton(rapidos_frame, label,
                        lambda c=clave: self._aplicar_filtro_rapido(c),
                        tipo='outline').pack(side='left', padx=3)

        # ========== CUERPO PRINCIPAL ==========
        body = tk.Frame(self.parent, bg=COLORES['fondo'])
        body.pack(fill='both', expand=True, padx=15, pady=5)

        # ========== SECCIÓN 1: KPIs PRINCIPALES ==========
        kpi_frame = tk.Frame(body, bg=COLORES['fondo'])
        kpi_frame.pack(fill='x', padx=0, pady=(0, 15))

        # Total Ventas
        card1 = tk.Frame(kpi_frame, bg=COLORES['fondo_card'],
                        highlightbackground=COLORES['borde'], highlightthickness=1)
        card1.pack(side='left', fill='both', expand=True, padx=(0, 8))

        tk.Label(card1, text="Total Ventas", font=FUENTES['pequena'],
                 fg=COLORES['texto_secundario'], bg=COLORES['fondo_card']).pack(anchor='w', padx=15, pady=(15, 5))
        self.lbl_total_ventas = tk.Label(card1, text="$ 0", font=FUENTES['kpi_valor'],
                                        fg=COLORES['primario'], bg=COLORES['fondo_card'])
        self.lbl_total_ventas.pack(anchor='w', padx=15, pady=(0, 15))

        # Total Boletas
        card2 = tk.Frame(kpi_frame, bg=COLORES['fondo_card'],
                        highlightbackground=COLORES['borde'], highlightthickness=1)
        card2.pack(side='left', fill='both', expand=True, padx=(0, 8))

        tk.Label(card2, text="Total Boletas", font=FUENTES['pequena'],
                 fg=COLORES['texto_secundario'], bg=COLORES['fondo_card']).pack(anchor='w', padx=15, pady=(15, 5))
        self.lbl_total_boletas = tk.Label(card2, text="$ 0", font=FUENTES['kpi_valor'],
                                         fg=COLORES['acento'], bg=COLORES['fondo_card'])
        self.lbl_total_boletas.pack(anchor='w', padx=15, pady=(0, 15))

        # Total Gastos
        card3 = tk.Frame(kpi_frame, bg=COLORES['fondo_card'],
                        highlightbackground=COLORES['borde'], highlightthickness=1)
        card3.pack(side='left', fill='both', expand=True, padx=(0, 8))

        tk.Label(card3, text="Total Gastos", font=FUENTES['pequena'],
                 fg=COLORES['texto_secundario'], bg=COLORES['fondo_card']).pack(anchor='w', padx=15, pady=(15, 5))
        self.lbl_total_gastos = tk.Label(card3, text="$ 0", font=FUENTES['kpi_valor'],
                                        fg=COLORES['error'], bg=COLORES['fondo_card'])
        self.lbl_total_gastos.pack(anchor='w', padx=15, pady=(0, 15))

        # Utilidad
        card4 = tk.Frame(kpi_frame, bg=COLORES['fondo_card'],
                        highlightbackground=COLORES['borde'], highlightthickness=1)
        card4.pack(side='left', fill='both', expand=True, padx=(0, 8))

        tk.Label(card4, text="Utilidad Neta", font=FUENTES['pequena'],
                 fg=COLORES['texto_secundario'], bg=COLORES['fondo_card']).pack(anchor='w', padx=15, pady=(15, 5))
        self.lbl_utilidad = tk.Label(card4, text="$ 0", font=FUENTES['kpi_valor'],
                                    fg=COLORES['exito'], bg=COLORES['fondo_card'])
        self.lbl_utilidad.pack(anchor='w', padx=15, pady=(0, 15))

        # Gasto promedio por cliente (boletas)
        card5 = tk.Frame(kpi_frame, bg=COLORES['fondo_card'],
                        highlightbackground=COLORES['borde'], highlightthickness=1)
        card5.pack(side='left', fill='both', expand=True)

        tk.Label(card5, text="Gasto Prom. / Cliente", font=FUENTES['pequena'],
                 fg=COLORES['texto_secundario'], bg=COLORES['fondo_card']).pack(anchor='w', padx=15, pady=(15, 5))
        tk.Label(card5, text="(boleta + consumos)", font=FUENTES['pequena'],
                 fg=COLORES['texto_deshabilitado'], bg=COLORES['fondo_card']).pack(anchor='w', padx=15)
        self.lbl_gasto_promedio = tk.Label(card5, text="$ 0", font=FUENTES['kpi_valor'],
                                           fg=COLORES['primario'], bg=COLORES['fondo_card'])
        self.lbl_gasto_promedio.pack(anchor='w', padx=15, pady=(0, 15))

        # ========== SECCIÓN 2: TABLAS ==========
        notebook = ttk.Notebook(body)
        notebook.pack(fill='both', expand=True)

        # Tab 1: Resumen de Categorías
        tab1 = tk.Frame(notebook, bg=COLORES['fondo'])
        notebook.add(tab1, text='Por Categoría')

        tk.Label(tab1, text="Ventas por Categoría", font=FUENTES['subtitulo'],
                 fg=COLORES['texto'], bg=COLORES['fondo']).pack(anchor='w', padx=15, pady=(10, 5))

        tabla_frame1 = tk.Frame(tab1, bg=COLORES['fondo_card'],
                               highlightbackground=COLORES['borde'], highlightthickness=1)
        tabla_frame1.pack(fill='both', expand=True, padx=15, pady=(0, 15))

        self.tree_categorias = ttk.Treeview(tabla_frame1, height=10,
                                            columns=('categoria', 'cantidad', 'total'),
                                            show='headings')

        self.tree_categorias.column('categoria', width=300, anchor='w')
        self.tree_categorias.column('cantidad', width=150, anchor='center')
        self.tree_categorias.column('total', width=200, anchor='e')

        self.tree_categorias.heading('categoria', text='Categoría')
        self.tree_categorias.heading('cantidad', text='Unidades Vendidas')
        self.tree_categorias.heading('total', text='Total')

        aplicar_estilo_tabla(self.tree_categorias)
        self.tree_categorias.pack(fill='both', expand=True)

        def _scroll_rep(event):
            d = int(-1 * (event.delta / 120)) if event.delta else (1 if event.num == 5 else -1)
            self.tree_categorias.yview_scroll(d, 'units')
        self.tree_categorias.bind_all('<MouseWheel>', _scroll_rep)
        self.tree_categorias.bind_all('<Button-4>', _scroll_rep)
        self.tree_categorias.bind_all('<Button-5>', _scroll_rep)

        # Tab 2: Top Productos
        tab2 = tk.Frame(notebook, bg=COLORES['fondo'])
        notebook.add(tab2, text='Top Productos')

        tk.Label(tab2, text="Top 10 Productos Más Vendidos", font=FUENTES['subtitulo'],
                 fg=COLORES['texto'], bg=COLORES['fondo']).pack(anchor='w', padx=15, pady=(10, 5))

        tabla_frame2 = tk.Frame(tab2, bg=COLORES['fondo_card'],
                               highlightbackground=COLORES['borde'], highlightthickness=1)
        tabla_frame2.pack(fill='both', expand=True, padx=15, pady=(0, 15))

        self.tree_productos = ttk.Treeview(tabla_frame2, height=10,
                                           columns=('producto', 'cantidad', 'precio_unit', 'total'),
                                           show='headings')

        self.tree_productos.column('producto', width=250, anchor='w')
        self.tree_productos.column('cantidad', width=120, anchor='center')
        self.tree_productos.column('precio_unit', width=150, anchor='e')
        self.tree_productos.column('total', width=150, anchor='e')

        self.tree_productos.heading('producto', text='Producto')
        self.tree_productos.heading('cantidad', text='Cantidad')
        self.tree_productos.heading('precio_unit', text='Precio Unit.')
        self.tree_productos.heading('total', text='Total')

        aplicar_estilo_tabla(self.tree_productos)
        self.tree_productos.pack(fill='both', expand=True)

        # Tab 3: Métodos de Pago
        tab3 = tk.Frame(notebook, bg=COLORES['fondo'])
        notebook.add(tab3, text='Métodos de Pago')

        tk.Label(tab3, text="Resumen por Método de Pago", font=FUENTES['subtitulo'],
                 fg=COLORES['texto'], bg=COLORES['fondo']).pack(anchor='w', padx=15, pady=(10, 5))

        tabla_frame3 = tk.Frame(tab3, bg=COLORES['fondo_card'],
                               highlightbackground=COLORES['borde'], highlightthickness=1)
        tabla_frame3.pack(fill='both', expand=True, padx=15, pady=(0, 15))

        self.tree_metodos = ttk.Treeview(tabla_frame3, height=10,
                                         columns=('metodo', 'cantidad', 'total'),
                                         show='headings')

        self.tree_metodos.column('metodo', width=300, anchor='w')
        self.tree_metodos.column('cantidad', width=150, anchor='center')
        self.tree_metodos.column('total', width=200, anchor='e')

        self.tree_metodos.heading('metodo', text='Método de Pago')
        self.tree_metodos.heading('cantidad', text='Transacciones')
        self.tree_metodos.heading('total', text='Total')

        aplicar_estilo_tabla(self.tree_metodos)
        self.tree_metodos.pack(fill='both', expand=True)

        # Tab 4: IVA 19%
        tab4 = tk.Frame(notebook, bg=COLORES['fondo'])
        notebook.add(tab4, text='IVA 19%')

        # Subtítulo y nota
        hdr4 = tk.Frame(tab4, bg=COLORES['fondo'])
        hdr4.pack(fill='x', padx=15, pady=(10, 0))
        tk.Label(hdr4, text="Reporte de IVA 19%", font=FUENTES['subtitulo'],
                 fg=COLORES['texto'], bg=COLORES['fondo']).pack(side='left')
        tk.Label(hdr4,
                 text="(Precios con IVA incluido — base = total / 1.19)",
                 font=FUENTES['pequena'],
                 fg=COLORES['texto_secundario'], bg=COLORES['fondo']).pack(side='left', padx=12)

        # KPIs IVA
        iva_kpi = tk.Frame(tab4, bg=COLORES['fondo'])
        iva_kpi.pack(fill='x', padx=15, pady=(8, 10))

        def _kpi_card(parent, titulo, attr_lbl, color):
            card = tk.Frame(parent, bg=COLORES['fondo_card'],
                            highlightbackground=COLORES['borde'], highlightthickness=1)
            card.pack(side='left', fill='both', expand=True, padx=(0, 6))
            tk.Frame(card, bg=color, height=3).pack(fill='x')
            tk.Label(card, text=titulo, font=FUENTES['pequena'],
                     fg=COLORES['texto_secundario'],
                     bg=COLORES['fondo_card']).pack(anchor='w', padx=12, pady=(8, 4))
            lbl = tk.Label(card, text="$ 0", font=FUENTES['kpi_valor'],
                           fg=color, bg=COLORES['fondo_card'])
            lbl.pack(anchor='w', padx=12, pady=(0, 10))
            setattr(self, attr_lbl, lbl)

        _kpi_card(iva_kpi, "Total con IVA", 'lbl_iva_total',    COLORES['primario'])
        _kpi_card(iva_kpi, "Base Gravable", 'lbl_iva_base',     COLORES['texto'])
        _kpi_card(iva_kpi, "IVA 19%",       'lbl_iva_valor',    COLORES['advertencia']
                  if 'advertencia' in COLORES else '#E65100')
        _kpi_card(iva_kpi, "Facturas Elect.", 'lbl_iva_fe',     COLORES['agua']
                  if 'agua' in COLORES else COLORES['primario'])

        # Tabla desglose por método de pago con IVA
        tk.Label(tab4, text="Desglose por Metodo de Pago",
                 font=FUENTES['normal_bold'], fg=COLORES['texto'],
                 bg=COLORES['fondo']).pack(anchor='w', padx=15, pady=(0, 4))

        tabla_frame4 = tk.Frame(tab4, bg=COLORES['fondo_card'],
                                highlightbackground=COLORES['borde'], highlightthickness=1)
        tabla_frame4.pack(fill='both', expand=True, padx=15, pady=(0, 15))

        self.tree_iva = ttk.Treeview(
            tabla_frame4, height=8,
            columns=('metodo', 'transacciones', 'total_iva', 'base', 'iva'),
            show='headings'
        )
        self.tree_iva.column('metodo',        width=200, anchor='w')
        self.tree_iva.column('transacciones', width=120, anchor='center')
        self.tree_iva.column('total_iva',     width=160, anchor='e')
        self.tree_iva.column('base',          width=160, anchor='e')
        self.tree_iva.column('iva',           width=130, anchor='e')

        self.tree_iva.heading('metodo',        text='Metodo de Pago')
        self.tree_iva.heading('transacciones', text='Transacciones')
        self.tree_iva.heading('total_iva',     text='Total con IVA')
        self.tree_iva.heading('base',          text='Base Gravable')
        self.tree_iva.heading('iva',           text='IVA 19%')

        aplicar_estilo_tabla(self.tree_iva)
        self.tree_iva.pack(fill='both', expand=True)

        # Tab 5: Anulaciones
        tab5 = tk.Frame(notebook, bg=COLORES['fondo'])
        notebook.add(tab5, text='Anulaciones')

        tk.Label(tab5, text="Ventas Anuladas", font=FUENTES['subtitulo'],
                 fg=COLORES['texto'], bg=COLORES['fondo']).pack(anchor='w', padx=15, pady=(10, 5))

        tabla_frame5 = tk.Frame(tab5, bg=COLORES['fondo_card'],
                                highlightbackground=COLORES['borde'], highlightthickness=1)
        tabla_frame5.pack(fill='both', expand=True, padx=15, pady=(0, 15))

        self.tree_anulaciones = ttk.Treeview(
            tabla_frame5, height=12,
            columns=('numero', 'fecha', 'cliente', 'total', 'usuario', 'motivo'),
            show='headings'
        )
        self.tree_anulaciones.column('numero',  width=140, anchor='w')
        self.tree_anulaciones.column('fecha',   width=130, anchor='center')
        self.tree_anulaciones.column('cliente', width=160, anchor='w')
        self.tree_anulaciones.column('total',   width=110, anchor='e')
        self.tree_anulaciones.column('usuario', width=110, anchor='center')
        self.tree_anulaciones.column('motivo',  width=250, anchor='w')

        self.tree_anulaciones.heading('numero',  text='N. Venta')
        self.tree_anulaciones.heading('fecha',   text='Fecha')
        self.tree_anulaciones.heading('cliente', text='Cliente')
        self.tree_anulaciones.heading('total',   text='Total')
        self.tree_anulaciones.heading('usuario', text='Usuario')
        self.tree_anulaciones.heading('motivo',  text='Motivo')

        aplicar_estilo_tabla(self.tree_anulaciones)
        self.tree_anulaciones.pack(fill='both', expand=True)

        # Tab 6: Por Cajero/Usuario
        tab6 = tk.Frame(notebook, bg=COLORES['fondo'])
        notebook.add(tab6, text='Por Cajero')

        tk.Label(tab6, text="Ventas por Cajero / Usuario", font=FUENTES['subtitulo'],
                 fg=COLORES['texto'], bg=COLORES['fondo']).pack(anchor='w', padx=15, pady=(10, 5))

        tabla_frame6 = tk.Frame(tab6, bg=COLORES['fondo_card'],
                                highlightbackground=COLORES['borde'], highlightthickness=1)
        tabla_frame6.pack(fill='both', expand=True, padx=15, pady=(0, 15))

        self.tree_cajeros = ttk.Treeview(
            tabla_frame6, height=12,
            columns=('usuario', 'transacciones', 'total', 'promedio', 'anulaciones'),
            show='headings'
        )
        self.tree_cajeros.column('usuario',       width=180, anchor='w')
        self.tree_cajeros.column('transacciones', width=120, anchor='center')
        self.tree_cajeros.column('total',         width=140, anchor='e')
        self.tree_cajeros.column('promedio',      width=130, anchor='e')
        self.tree_cajeros.column('anulaciones',   width=110, anchor='center')

        self.tree_cajeros.heading('usuario',       text='Cajero / Usuario')
        self.tree_cajeros.heading('transacciones', text='Ventas')
        self.tree_cajeros.heading('total',         text='Total Vendido')
        self.tree_cajeros.heading('promedio',      text='Ticket Promedio')
        self.tree_cajeros.heading('anulaciones',   text='Anulaciones')

        aplicar_estilo_tabla(self.tree_cajeros)
        self.tree_cajeros.pack(fill='both', expand=True)

        # Tab 7: Sesgo / Anomalías
        tab7 = tk.Frame(notebook, bg=COLORES['fondo'])
        notebook.add(tab7, text='Sesgo / Anomalias')

        tk.Label(tab7, text="Anulaciones por cajero",
                 font=FUENTES['subtitulo'], bg=COLORES['fondo'],
                 fg=COLORES['texto']).pack(anchor='w', padx=10, pady=(10, 2))

        cols_an = ('usuario', 'ventas_ok', 'anulaciones', 'tasa_pct', 'total_anulado')
        self.tree_sesgo_anulaciones = ttk.Treeview(tab7, columns=cols_an,
                                                   show='headings', height=5)
        self.tree_sesgo_anulaciones.column('usuario',       width=170, anchor='w')
        self.tree_sesgo_anulaciones.column('ventas_ok',     width=90,  anchor='center')
        self.tree_sesgo_anulaciones.column('anulaciones',   width=90,  anchor='center')
        self.tree_sesgo_anulaciones.column('tasa_pct',      width=100, anchor='center')
        self.tree_sesgo_anulaciones.column('total_anulado', width=140, anchor='e')
        self.tree_sesgo_anulaciones.heading('usuario',       text='Cajero')
        self.tree_sesgo_anulaciones.heading('ventas_ok',     text='Ventas OK')
        self.tree_sesgo_anulaciones.heading('anulaciones',   text='Anulaciones')
        self.tree_sesgo_anulaciones.heading('tasa_pct',      text='Tasa (%)')
        self.tree_sesgo_anulaciones.heading('total_anulado', text='Total Anulado')
        aplicar_estilo_tabla(self.tree_sesgo_anulaciones)
        self.tree_sesgo_anulaciones.tag_configure('alerta', background='#FFF3CD')
        self.tree_sesgo_anulaciones.pack(fill='x', padx=10)

        tk.Label(tab7, text="Descuentos por cajero",
                 font=FUENTES['subtitulo'], bg=COLORES['fondo'],
                 fg=COLORES['texto']).pack(anchor='w', padx=10, pady=(12, 2))

        cols_desc = ('usuario', 'ventas_con_desc', 'total_descuento', 'desc_promedio', 'max_descuento')
        self.tree_sesgo_descuentos = ttk.Treeview(tab7, columns=cols_desc,
                                                  show='headings', height=5)
        self.tree_sesgo_descuentos.column('usuario',         width=170, anchor='w')
        self.tree_sesgo_descuentos.column('ventas_con_desc', width=120, anchor='center')
        self.tree_sesgo_descuentos.column('total_descuento', width=140, anchor='e')
        self.tree_sesgo_descuentos.column('desc_promedio',   width=140, anchor='e')
        self.tree_sesgo_descuentos.column('max_descuento',   width=140, anchor='e')
        self.tree_sesgo_descuentos.heading('usuario',         text='Cajero')
        self.tree_sesgo_descuentos.heading('ventas_con_desc', text='Ventas c/Descuento')
        self.tree_sesgo_descuentos.heading('total_descuento', text='Total Descontado')
        self.tree_sesgo_descuentos.heading('desc_promedio',   text='Descuento Prom.')
        self.tree_sesgo_descuentos.heading('max_descuento',   text='Max Descuento')
        aplicar_estilo_tabla(self.tree_sesgo_descuentos)
        self.tree_sesgo_descuentos.pack(fill='x', padx=10)

        tk.Label(tab7, text="Gastos por registrador",
                 font=FUENTES['subtitulo'], bg=COLORES['fondo'],
                 fg=COLORES['texto']).pack(anchor='w', padx=10, pady=(12, 2))

        cols_gas = ('usuario', 'num_gastos', 'total_gastos', 'prom_gasto', 'max_gasto')
        self.tree_sesgo_gastos = ttk.Treeview(tab7, columns=cols_gas,
                                              show='headings', height=5)
        self.tree_sesgo_gastos.column('usuario',    width=170, anchor='w')
        self.tree_sesgo_gastos.column('num_gastos', width=90,  anchor='center')
        self.tree_sesgo_gastos.column('total_gastos', width=140, anchor='e')
        self.tree_sesgo_gastos.column('prom_gasto', width=140, anchor='e')
        self.tree_sesgo_gastos.column('max_gasto',  width=140, anchor='e')
        self.tree_sesgo_gastos.heading('usuario',    text='Registrador')
        self.tree_sesgo_gastos.heading('num_gastos', text='N. Gastos')
        self.tree_sesgo_gastos.heading('total_gastos', text='Total Gastos')
        self.tree_sesgo_gastos.heading('prom_gasto', text='Promedio')
        self.tree_sesgo_gastos.heading('max_gasto',  text='Mayor Gasto')
        aplicar_estilo_tabla(self.tree_sesgo_gastos)
        self.tree_sesgo_gastos.pack(fill='x', padx=10)

    def _aplicar_filtro_rapido(self, periodo):
        """Aplica un filtro de fecha predefinido y regenera los reportes"""
        hoy = datetime.date.today()
        if periodo == 'hoy':
            inicio = fin = hoy
        elif periodo == 'semana':
            inicio = hoy - datetime.timedelta(days=hoy.weekday())  # Lunes
            fin = hoy
        elif periodo == 'mes':
            inicio = hoy.replace(day=1)
            fin = hoy
        elif periodo == 'mes_anterior':
            fin = hoy.replace(day=1) - datetime.timedelta(days=1)
            inicio = fin.replace(day=1)
        else:
            return

        self.entry_desde.delete(0, tk.END)
        self.entry_desde.insert(0, inicio.isoformat())
        self.entry_hasta.delete(0, tk.END)
        self.entry_hasta.insert(0, fin.isoformat())
        self._generar_reportes()

    def _generar_reportes(self, event=None):
        """Genera todos los reportes"""
        try:
            fecha_inicio = self.entry_desde.get()
            fecha_fin = self.entry_hasta.get()

            with conexion_segura() as conn:
                # ========== KPIs ==========
                venta_data = conn.execute("""
                    SELECT COALESCE(SUM(total), 0) as total, COUNT(*) as cantidad
                    FROM ventas
                    WHERE tipo IN ('normal', 'cuenta_abierta', 'para_llevar')
                      AND estado != 'ANULADA'
                      AND DATE(fecha_creacion) BETWEEN ? AND ?
                """, (fecha_inicio, fecha_fin)).fetchone()

                boleta_data = conn.execute("""
                    SELECT COALESCE(SUM(total), 0) as total, COUNT(*) as cantidad
                    FROM boletas_entrada
                    WHERE DATE(hora_entrada) BETWEEN ? AND ?
                """, (fecha_inicio, fecha_fin)).fetchone()

                gasto_data = conn.execute("""
                    SELECT COALESCE(SUM(valor), 0) as total, COUNT(*) as cantidad
                    FROM gastos
                    WHERE DATE(fecha) BETWEEN ? AND ?
                """, (fecha_inicio, fecha_fin)).fetchone()

                total_ventas = venta_data['total']
                total_boletas = boleta_data['total']
                total_gastos = gasto_data['total']
                utilidad = total_ventas + total_boletas - total_gastos

                self.lbl_total_ventas.config(text=format_money(total_ventas))
                self.lbl_total_boletas.config(text=format_money(total_boletas))
                self.lbl_total_gastos.config(text=format_money(total_gastos))
                self.lbl_utilidad.config(text=format_money(utilidad))

                # Gasto promedio por cliente = (total_boletas + total_ventas) / personas_en_boletas
                personas_data = conn.execute("""
                    SELECT COALESCE(SUM(cantidad_personas), 0) as personas
                    FROM boletas_entrada
                    WHERE DATE(hora_entrada) BETWEEN ? AND ?
                """, (fecha_inicio, fecha_fin)).fetchone()
                personas = personas_data['personas'] if personas_data else 0
                if personas and personas > 0:
                    gasto_prom = round((total_boletas + total_ventas) / personas, 0)
                else:
                    gasto_prom = 0
                self.lbl_gasto_promedio.config(text=format_money(gasto_prom))

                # ========== VENTAS POR CATEGORÍA ==========
                categorias = conn.execute("""
                    SELECT c.nombre, SUM(vd.cantidad) as cantidad, SUM(vd.total_linea) as total
                    FROM venta_detalle vd
                    JOIN productos p ON vd.id_producto = p.id_producto
                    LEFT JOIN categorias c ON p.id_categoria = c.id_categoria
                    JOIN ventas v ON vd.id_venta = v.id_venta
                    WHERE DATE(v.fecha_creacion) BETWEEN ? AND ?
                    GROUP BY c.nombre
                    ORDER BY total DESC
                """, (fecha_inicio, fecha_fin)).fetchall()

                for item in self.tree_categorias.get_children():
                    self.tree_categorias.delete(item)

                for i, cat in enumerate(categorias):
                    tag = 'par' if i % 2 == 0 else 'impar'
                    self.tree_categorias.insert('', 'end', values=(
                        cat['nombre'] or 'Sin Categoría',
                        f"{cat['cantidad']} unidades",
                        format_money(cat['total'])
                    ), tags=(tag,))

                # ========== TOP PRODUCTOS ==========
                productos = conn.execute("""
                    SELECT p.nombre, SUM(vd.cantidad) as cantidad, vd.precio_unitario, SUM(vd.total_linea) as total
                    FROM venta_detalle vd
                    JOIN productos p ON vd.id_producto = p.id_producto
                    JOIN ventas v ON vd.id_venta = v.id_venta
                    WHERE DATE(v.fecha_creacion) BETWEEN ? AND ?
                    GROUP BY vd.id_producto
                    ORDER BY total DESC
                    LIMIT 10
                """, (fecha_inicio, fecha_fin)).fetchall()

                for item in self.tree_productos.get_children():
                    self.tree_productos.delete(item)

                for i, prod in enumerate(productos):
                    tag = 'par' if i % 2 == 0 else 'impar'
                    self.tree_productos.insert('', 'end', values=(
                        prod['nombre'],
                        f"{prod['cantidad']} unidades",
                        format_money(prod['precio_unitario']),
                        format_money(prod['total'])
                    ), tags=(tag,))

                # ========== MÉTODOS DE PAGO ==========
                metodos = conn.execute("""
                    SELECT metodo_pago, COUNT(*) as cantidad, SUM(total) as total
                    FROM ventas
                    WHERE DATE(fecha_creacion) BETWEEN ? AND ? AND estado IN ('CERRADA', 'PAGADA')
                    GROUP BY metodo_pago
                    ORDER BY total DESC
                """, (fecha_inicio, fecha_fin)).fetchall()

                for item in self.tree_metodos.get_children():
                    self.tree_metodos.delete(item)

                for i, met in enumerate(metodos):
                    tag = 'par' if i % 2 == 0 else 'impar'
                    self.tree_metodos.insert('', 'end', values=(
                        met['metodo_pago'] or 'No especificado',
                        met['cantidad'],
                        format_money(met['total'])
                    ), tags=(tag,))

                # ========== IVA 19% ==========
                # Ventas del período (excluye anuladas); precios incluyen IVA
                iva_metodos = conn.execute("""
                    SELECT metodo_pago,
                           COUNT(*) as cantidad,
                           COALESCE(SUM(total), 0) as total_con_iva
                    FROM ventas
                    WHERE tipo IN ('normal', 'cuenta_abierta')
                      AND estado IN ('CERRADA', 'PAGADA')
                      AND DATE(fecha_creacion) BETWEEN ? AND ?
                    GROUP BY metodo_pago
                    ORDER BY total_con_iva DESC
                """, (fecha_inicio, fecha_fin)).fetchall()

                total_con_iva_global = sum(r['total_con_iva'] for r in iva_metodos)
                base_global = round(total_con_iva_global / 1.19, 2)
                iva_global  = round(total_con_iva_global - base_global, 2)

                # Facturas electrónicas en el período
                fe_count = conn.execute("""
                    SELECT COUNT(*) as n FROM facturas_electronicas fe
                    JOIN ventas v ON fe.id_venta = v.id_venta
                    WHERE DATE(v.fecha_creacion) BETWEEN ? AND ?
                """, (fecha_inicio, fecha_fin)).fetchone()['n']

                self.lbl_iva_total.config(text=format_money(total_con_iva_global))
                self.lbl_iva_base.config(text=format_money(base_global))
                self.lbl_iva_valor.config(text=format_money(iva_global))
                self.lbl_iva_fe.config(text=str(fe_count))

                for item in self.tree_iva.get_children():
                    self.tree_iva.delete(item)

                for i, row in enumerate(iva_metodos):
                    tiva  = row['total_con_iva']
                    base  = round(tiva / 1.19, 2)
                    iva_f = round(tiva - base, 2)
                    tag   = 'par' if i % 2 == 0 else 'impar'
                    self.tree_iva.insert('', 'end', values=(
                        row['metodo_pago'] or 'No especificado',
                        row['cantidad'],
                        format_money(tiva),
                        format_money(base),
                        format_money(iva_f),
                    ), tags=(tag,))

                # Fila totales
                if iva_metodos:
                    self.tree_iva.insert('', 'end', values=(
                        'TOTAL',
                        sum(r['cantidad'] for r in iva_metodos),
                        format_money(total_con_iva_global),
                        format_money(base_global),
                        format_money(iva_global),
                    ), tags=('total',))
                    self.tree_iva.tag_configure('total',
                        background=COLORES['primario'],
                        foreground=COLORES['texto_claro'])

                # ========== ANULACIONES ==========
                anuladas = conn.execute("""
                    SELECT v.numero_venta, v.fecha_creacion, v.cliente_nombre,
                           v.total, v.usuario_nombre,
                           COALESCE(
                               (SELECT a.comentario FROM auditoria a
                                WHERE a.tabla_afectada = 'ventas' AND a.id_registro = v.id_venta
                                  AND a.accion = 'ANULAR'
                                ORDER BY a.fecha_hora DESC LIMIT 1),
                               'Sin motivo registrado'
                           ) as motivo
                    FROM ventas v
                    WHERE v.estado = 'ANULADA'
                      AND DATE(v.fecha_creacion) BETWEEN ? AND ?
                    ORDER BY v.fecha_creacion DESC
                """, (fecha_inicio, fecha_fin)).fetchall()

                for item in self.tree_anulaciones.get_children():
                    self.tree_anulaciones.delete(item)

                for i, an in enumerate(anuladas):
                    tag = 'par' if i % 2 == 0 else 'impar'
                    fecha_fmt = an['fecha_creacion'][:16] if an['fecha_creacion'] else ''
                    self.tree_anulaciones.insert('', 'end', values=(
                        an['numero_venta'],
                        fecha_fmt,
                        an['cliente_nombre'] or '-',
                        format_money(an['total']),
                        an['usuario_nombre'] or '-',
                        an['motivo'],
                    ), tags=(tag,))

                # ========== POR CAJERO / USUARIO ==========
                cajeros = conn.execute("""
                    SELECT usuario_nombre,
                           COUNT(CASE WHEN estado IN ('CERRADA','PAGADA') THEN 1 END) as transacciones,
                           COALESCE(SUM(CASE WHEN estado IN ('CERRADA','PAGADA') THEN total ELSE 0 END), 0) as total,
                           COUNT(CASE WHEN estado = 'ANULADA' THEN 1 END) as anulaciones
                    FROM ventas
                    WHERE tipo IN ('normal','cuenta_abierta','para_llevar')
                      AND DATE(fecha_creacion) BETWEEN ? AND ?
                    GROUP BY usuario_nombre
                    ORDER BY total DESC
                """, (fecha_inicio, fecha_fin)).fetchall()

                for item in self.tree_cajeros.get_children():
                    self.tree_cajeros.delete(item)

                for i, caj in enumerate(cajeros):
                    transac = caj['transacciones'] or 0
                    total_c = caj['total'] or 0
                    promedio = round(total_c / transac, 0) if transac else 0
                    tag = 'par' if i % 2 == 0 else 'impar'
                    self.tree_cajeros.insert('', 'end', values=(
                        caj['usuario_nombre'] or 'Sin nombre',
                        transac,
                        format_money(total_c),
                        format_money(promedio),
                        caj['anulaciones'],
                    ), tags=(tag,))

                # ========== SESGO / ANULACIONES POR CAJERO ==========
                sesgo_anulaciones = conn.execute("""
                    SELECT usuario_nombre,
                           COUNT(CASE WHEN estado != 'ANULADA' THEN 1 END) as ventas_ok,
                           COUNT(CASE WHEN estado = 'ANULADA' THEN 1 END) as anuladas,
                           ROUND(
                               100.0 * COUNT(CASE WHEN estado = 'ANULADA' THEN 1 END)
                               / NULLIF(COUNT(*), 0), 1
                           ) as tasa_anulacion,
                           COALESCE(SUM(CASE WHEN estado = 'ANULADA' THEN total ELSE 0 END), 0) as total_anulado
                    FROM ventas
                    WHERE tipo IN ('normal','cuenta_abierta','para_llevar','boleta')
                      AND DATE(fecha_creacion) BETWEEN ? AND ?
                    GROUP BY usuario_nombre
                    HAVING (ventas_ok + anuladas) > 0
                    ORDER BY tasa_anulacion DESC
                """, (fecha_inicio, fecha_fin)).fetchall()

                for item in self.tree_sesgo_anulaciones.get_children():
                    self.tree_sesgo_anulaciones.delete(item)

                for row in sesgo_anulaciones:
                    tasa = row['tasa_anulacion'] or 0.0
                    tags = ('alerta',) if tasa > 5.0 else ()
                    self.tree_sesgo_anulaciones.insert('', 'end', values=(
                        row['usuario_nombre'] or 'Sin nombre',
                        row['ventas_ok'],
                        row['anuladas'],
                        f"{tasa:.1f}%",
                        format_money(row['total_anulado']),
                    ), tags=tags)

                # ========== SESGO / DESCUENTOS POR CAJERO ==========
                sesgo_descuentos = conn.execute("""
                    SELECT usuario_nombre,
                           COUNT(*) as ventas_con_descuento,
                           COALESCE(SUM(descuento), 0) as total_descuentado,
                           ROUND(AVG(descuento), 0) as promedio_descuento,
                           COALESCE(MAX(descuento), 0) as max_descuento
                    FROM ventas
                    WHERE descuento > 0
                      AND estado != 'ANULADA'
                      AND DATE(fecha_creacion) BETWEEN ? AND ?
                    GROUP BY usuario_nombre
                    ORDER BY total_descuentado DESC
                """, (fecha_inicio, fecha_fin)).fetchall()

                for item in self.tree_sesgo_descuentos.get_children():
                    self.tree_sesgo_descuentos.delete(item)

                for i, row in enumerate(sesgo_descuentos):
                    tag = 'par' if i % 2 == 0 else 'impar'
                    self.tree_sesgo_descuentos.insert('', 'end', values=(
                        row['usuario_nombre'] or 'Sin nombre',
                        row['ventas_con_descuento'],
                        format_money(row['total_descuentado']),
                        format_money(row['promedio_descuento']),
                        format_money(row['max_descuento']),
                    ), tags=(tag,))

                # ========== SESGO / GASTOS POR REGISTRADOR ==========
                sesgo_gastos = conn.execute("""
                    SELECT usuario_registro,
                           COUNT(*) as num_gastos,
                           COALESCE(SUM(valor), 0) as total_gastado,
                           ROUND(AVG(valor), 0) as promedio_gasto,
                           COALESCE(MAX(valor), 0) as max_gasto
                    FROM gastos
                    WHERE DATE(fecha) BETWEEN ? AND ?
                    GROUP BY usuario_registro
                    ORDER BY total_gastado DESC
                """, (fecha_inicio, fecha_fin)).fetchall()

                for item in self.tree_sesgo_gastos.get_children():
                    self.tree_sesgo_gastos.delete(item)

                for i, row in enumerate(sesgo_gastos):
                    tag = 'par' if i % 2 == 0 else 'impar'
                    self.tree_sesgo_gastos.insert('', 'end', values=(
                        row['usuario_registro'] or 'Sin nombre',
                        row['num_gastos'],
                        format_money(row['total_gastado']),
                        format_money(row['promedio_gasto']),
                        format_money(row['max_gasto']),
                    ), tags=(tag,))

        except Exception as e:
            messagebox.showerror("Error", f"Error al generar reportes: {str(e)}")

    def _exportar_excel(self):
        """Exporta los reportes actuales a un archivo Excel (.xlsx)."""
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            messagebox.showerror(
                "Dependencia faltante",
                "Se necesita instalar openpyxl.\nEjecute: pip install openpyxl"
            )
            return

        fecha_inicio = self.entry_desde.get()
        fecha_fin = self.entry_hasta.get()

        wb = openpyxl.Workbook()

        # ---------- Estilos ----------
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill("solid", fgColor="1B5E20")  # Verde oscuro
        alt_fill    = PatternFill("solid", fgColor="E8F5E9")  # Verde claro alternado
        center_align = Alignment(horizontal="center", vertical="center")
        right_align  = Alignment(horizontal="right",  vertical="center")
        left_align   = Alignment(horizontal="left",   vertical="center")
        thin_side    = Side(style="thin", color="BDBDBD")
        thin_border  = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        title_font   = Font(bold=True, size=14, color="1B5E20")

        def estilo_header(ws, fila, ncols):
            for col in range(1, ncols + 1):
                cell = ws.cell(row=fila, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = thin_border

        def estilo_fila(ws, fila, ncols, par):
            for col in range(1, ncols + 1):
                cell = ws.cell(row=fila, column=col)
                if par:
                    cell.fill = alt_fill
                cell.border = thin_border

        def autofit(ws):
            for col in ws.columns:
                max_len = max((len(str(c.value or '')) for c in col), default=10)
                ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

        # ---------- Hoja 1: Resumen ----------
        ws_resumen = wb.active
        ws_resumen.title = "Resumen"

        ws_resumen["A1"] = "Club Los Pocitos Azufrados - Reporte de Ventas"
        ws_resumen["A1"].font = title_font
        ws_resumen["A2"] = f"Período: {fecha_inicio} al {fecha_fin}"
        ws_resumen["A2"].font = Font(italic=True, color="616161")

        ws_resumen.append([])
        ws_resumen.append(["Concepto", "Valor"])
        estilo_header(ws_resumen, 4, 2)
        resumen_data = [
            ("Total Ventas",  self.lbl_total_ventas.cget("text")),
            ("Total Boletas", self.lbl_total_boletas.cget("text")),
            ("Total Gastos",  self.lbl_total_gastos.cget("text")),
            ("Utilidad Neta", self.lbl_utilidad.cget("text")),
        ]
        for i, (concepto, valor) in enumerate(resumen_data):
            fila_n = 5 + i
            ws_resumen.cell(row=fila_n, column=1, value=concepto).alignment = left_align
            ws_resumen.cell(row=fila_n, column=2, value=valor).alignment = right_align
            estilo_fila(ws_resumen, fila_n, 2, i % 2 == 0)
        autofit(ws_resumen)

        # ---------- Hoja 2: Por Categoría ----------
        ws_cat = wb.create_sheet("Por Categoría")
        ws_cat.append(["Categoría", "Unidades Vendidas", "Total"])
        estilo_header(ws_cat, 1, 3)
        for i, item in enumerate(self.tree_categorias.get_children()):
            vals = self.tree_categorias.item(item)['values']
            fila_n = 2 + i
            for col, val in enumerate(vals, 1):
                ws_cat.cell(row=fila_n, column=col, value=val).alignment = (
                    right_align if col == 3 else (center_align if col == 2 else left_align)
                )
            estilo_fila(ws_cat, fila_n, 3, i % 2 == 0)
        autofit(ws_cat)

        # ---------- Hoja 3: Top Productos ----------
        ws_prod = wb.create_sheet("Top Productos")
        ws_prod.append(["Producto", "Cantidad", "Precio Unit.", "Total"])
        estilo_header(ws_prod, 1, 4)
        for i, item in enumerate(self.tree_productos.get_children()):
            vals = self.tree_productos.item(item)['values']
            fila_n = 2 + i
            for col, val in enumerate(vals, 1):
                ws_prod.cell(row=fila_n, column=col, value=val).alignment = (
                    right_align if col in (3, 4) else (center_align if col == 2 else left_align)
                )
            estilo_fila(ws_prod, fila_n, 4, i % 2 == 0)
        autofit(ws_prod)

        # ---------- Hoja 4: Métodos de Pago ----------
        ws_met = wb.create_sheet("Métodos de Pago")
        ws_met.append(["Método de Pago", "Transacciones", "Total"])
        estilo_header(ws_met, 1, 3)
        for i, item in enumerate(self.tree_metodos.get_children()):
            vals = self.tree_metodos.item(item)['values']
            fila_n = 2 + i
            for col, val in enumerate(vals, 1):
                ws_met.cell(row=fila_n, column=col, value=val).alignment = (
                    right_align if col == 3 else (center_align if col == 2 else left_align)
                )
            estilo_fila(ws_met, fila_n, 3, i % 2 == 0)
        autofit(ws_met)

        # ---------- Hoja 5: Detalle de Ventas ----------
        try:
            with conexion_segura() as conn:
                ventas_det = conn.execute("""
                    SELECT v.numero_venta, v.fecha_creacion, v.tipo, v.metodo_pago,
                           v.total, v.estado,
                           COALESCE(u.nombre_completo, v.usuario_nombre) as cajero
                    FROM ventas v
                    LEFT JOIN usuarios u ON u.id_usuario = v.id_usuario
                    WHERE tipo IN ('normal', 'cuenta_abierta', 'para_llevar')
                      AND estado != 'ANULADA'
                      AND DATE(v.fecha_creacion) BETWEEN ? AND ?
                    ORDER BY v.fecha_creacion
                """, (fecha_inicio, fecha_fin)).fetchall()

                gastos_det = conn.execute("""
                    SELECT fecha, descripcion, categoria, valor, usuario_registro
                    FROM gastos
                    WHERE DATE(fecha) BETWEEN ? AND ?
                    ORDER BY fecha
                """, (fecha_inicio, fecha_fin)).fetchall()
        except Exception:
            ventas_det = []
            gastos_det = []

        ws_det = wb.create_sheet("Ventas Detalle")
        ws_det.append(["N. Venta", "Fecha", "Hora", "Tipo", "Metodo Pago", "Cajero", "Total", "Estado"])
        estilo_header(ws_det, 1, 8)
        for i, r in enumerate(ventas_det):
            fecha_str = r['fecha_creacion'][:10] if r['fecha_creacion'] else ''
            hora_str  = r['fecha_creacion'][11:16] if r['fecha_creacion'] else ''
            fila_n = 2 + i
            ws_det.cell(row=fila_n, column=1, value=r['numero_venta']).alignment = left_align
            ws_det.cell(row=fila_n, column=2, value=fecha_str).alignment = center_align
            ws_det.cell(row=fila_n, column=3, value=hora_str).alignment = center_align
            ws_det.cell(row=fila_n, column=4, value=r['tipo']).alignment = center_align
            ws_det.cell(row=fila_n, column=5, value=r['metodo_pago'] or '-').alignment = center_align
            ws_det.cell(row=fila_n, column=6, value=r['cajero'] or '-').alignment = left_align
            ws_det.cell(row=fila_n, column=7, value=r['total']).alignment = right_align
            ws_det.cell(row=fila_n, column=8, value=r['estado']).alignment = center_align
            estilo_fila(ws_det, fila_n, 8, i % 2 == 0)
        autofit(ws_det)

        ws_gas = wb.create_sheet("Gastos Detalle")
        ws_gas.append(["Fecha", "Descripcion", "Categoria", "Registrado por", "Valor"])
        estilo_header(ws_gas, 1, 5)
        for i, r in enumerate(gastos_det):
            fila_n = 2 + i
            ws_gas.cell(row=fila_n, column=1, value=r['fecha']).alignment = center_align
            ws_gas.cell(row=fila_n, column=2, value=r['descripcion']).alignment = left_align
            ws_gas.cell(row=fila_n, column=3, value=r['categoria'] or '-').alignment = left_align
            ws_gas.cell(row=fila_n, column=4, value=r['usuario_registro'] or '-').alignment = left_align
            ws_gas.cell(row=fila_n, column=5, value=r['valor']).alignment = right_align
            estilo_fila(ws_gas, fila_n, 5, i % 2 == 0)
        autofit(ws_gas)

        # ---------- Guardar ----------
        nombre_archivo = f"reporte_pocitos_{fecha_inicio}_a_{fecha_fin}.xlsx"
        ruta = os.path.join(BASE_DIR, 'data', nombre_archivo)
        wb.save(ruta)

        messagebox.showinfo(
            "Exportado",
            f"Reporte guardado en:\n{ruta}\n\n"
            f"Hojas incluidas:\n"
            f"  - Resumen\n  - Por Categoria\n  - Top Productos\n"
            f"  - Metodos de Pago\n  - Ventas Detalle\n  - Gastos Detalle"
        )

        # Intentar abrir automáticamente
        try:
            import subprocess
            subprocess.Popen(['start', '', ruta], shell=True)
        except Exception:
            pass

    def detener(self):
        try:
            self.tree_categorias.unbind_all('<MouseWheel>')
            self.tree_categorias.unbind_all('<Button-4>')
            self.tree_categorias.unbind_all('<Button-5>')
        except Exception:
            pass
